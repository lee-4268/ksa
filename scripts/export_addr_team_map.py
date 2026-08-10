#!/usr/bin/env python3
"""주소 → 품질개선팀 매핑 로직을 xlsx로 내보낸다.

매핑 규칙(상수/함수)은 yolov8/api/routers/inspection.py 소스에서 직접 추출해 실행하므로
코드가 바뀌면 결과물도 자동으로 따라간다(별도 복사본 유지 불필요).

사용법:
  python scripts/export_addr_team_map.py                       # 서울 확정규칙만 반영
  python scripts/export_addr_team_map.py -l /tmp/learned_addr_map.json
  python scripts/export_addr_team_map.py -l /tmp/cert_cache.db   # EC2에서 직접 학습
  python scripts/export_addr_team_map.py -o docs/주소-팀_매핑.xlsx --with-ri
"""
import argparse
import ast
import gzip
import json
import os
import sqlite3
import sys
import tempfile
import types

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INSPECTION_PY = ''
LEGAL_DONG_TSV = ''

# 행정동↔법정동 매핑 (행정안전부 '행정기관(행정동) 및 관할구역(법정동)' KIKmix, 2026.3.25 시행)
# 출처: https://www.mois.go.kr → 업무안내 → 주민등록,인감 → 변경내역 알림 첨부 jscode*.zip
# 이용허락범위 제한 없음. 갱신하려면 KIKmix xlsx 를 받아 --admin-dong 으로 넘기거나 이 파일을 교체.
ADMIN_DONG_TSV_GZ = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                 'data', 'admin_dong_map.tsv.gz')

# 품질개선팀(SKO) → SKT Access운용팀. 출처: 사내 SKTSKO조직맵핑.csv
# 원본에 남양주품질개선팀이 경기/인천 두 줄로 있어 인천Access운용팀으로 확정했다.
SKT_OPS_TSV = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           'data', 'skt_ops_team_map.tsv')

# 로컬 리포지토리 / EC2 배포본(APP_DIR) 양쪽에서 동작
API_DIR_CANDIDATES = [
    os.path.join(ROOT, 'yolov8', 'api'),
    os.environ.get('APP_DIR', '/home/ubuntu/kca-api'),
]


def resolve_paths(api_dir):
    """inspection.py 와 legal_dong_code.tsv 를 각각 탐색.

    배포본에는 deploy_backend.sh 가 코드만 복사하고 tsv 는 런타임 자산으로 남아 있어
    두 파일이 다른 위치에 있을 수 있으므로 따로 찾는다.
    """
    global INSPECTION_PY, LEGAL_DONG_TSV
    # --api-dir 를 명시했으면 폴백하지 않는다 (오타를 조용히 넘기지 않도록)
    cands = [api_dir] if api_dir else API_DIR_CANDIDATES

    for d in cands:
        p = os.path.join(d, 'routers', 'inspection.py')
        if os.path.exists(p):
            INSPECTION_PY = p
            break
    else:
        sys.exit('inspection.py 를 찾지 못했습니다. --api-dir 로 지정하세요.\n'
                 f'  탐색: {", ".join(cands)}')

    for d in cands:
        p = os.path.join(d, 'legal_dong_code.tsv')
        if os.path.exists(p):
            LEGAL_DONG_TSV = p
            break
    else:
        sys.exit('legal_dong_code.tsv 를 찾지 못했습니다.\n'
                 f'  탐색: {", ".join(cands)}\n'
                 '  저장소 yolov8/api/legal_dong_code.tsv 를 해당 위치에 두거나 '
                 '--api-dir 로 지정하세요.')

    return os.path.dirname(os.path.dirname(INSPECTION_PY))

WANT_CONSTS = [
    'INSP_ORG_MAP', 'INSP_TEAM_TO_HDQT', '_VALID_SKT_HDQTS', '_ACCESS_TO_SKT_HDQT',
    '_DEPRECATED_TEAM_MAP', '_SEOUL_GU_TO_TEAM', '_SEOUL_GU_SORTED', '_ADDR_ABBR_MAP',
    '_NON_GEOGRAPHIC_TEAMS',
]
WANT_FUNCS = ['_normalize_addr', '_hdqt_from_addr', '_normalize_skt_hdqt',
              '_normalize_learned_map', '_learn_addr_map_from_cert_db']


def load_logic():
    """inspection.py에서 매핑 상수/함수만 뽑아 격리된 네임스페이스에 적재."""
    src = open(INSPECTION_PY, encoding='utf-8').read()
    tree = ast.parse(src)
    lines = src.splitlines(keepends=True)

    def seg(node):
        return ''.join(lines[node.lineno - 1:node.end_lineno])

    chunks = []
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in WANT_FUNCS:
            chunks.append(seg(node))
        elif isinstance(node, (ast.Assign, ast.AnnAssign)):
            targets = node.targets if isinstance(node, ast.Assign) else [node.target]
            names = [t.id for t in targets if isinstance(t, ast.Name)]
            if any(n in WANT_CONSTS for n in names):
                chunks.append(seg(node))

    ns = {
        'os': os, 'sqlite3': sqlite3,
        # 학습 진행 상황을 그대로 보여준다 (제외 건수 등 확인용)
        'logger': types.SimpleNamespace(warning=print, info=print, error=print),
        '_cert_cache_mod': types.SimpleNamespace(_cert_cache_db_path=''),
    }
    exec(''.join(chunks), ns)

    missing = [n for n in WANT_CONSTS + WANT_FUNCS if n not in ns]
    if missing:
        sys.exit(f'inspection.py에서 다음을 찾지 못했습니다: {", ".join(missing)}\n'
                 f'  대상 파일: {INSPECTION_PY}\n'
                 '  이 스크립트보다 배포된 백엔드 코드가 오래되면 발생합니다.\n'
                 '  EC2 라면 코드부터 갱신하세요 (재시작 불필요):\n'
                 '      bash /home/ubuntu/deploy_backend.sh')
    return ns


class RecordingMap(dict):
    """_hdqt_from_addr가 어떤 키워드로 팀을 찾았는지 기록."""
    hit = None

    def get(self, key, default=None):
        val = super().get(key, default)
        if val:
            self.hit = key
        return val


def load_learned(path, logic):
    """학습맵 로드. 경로 미지정 시 서버 캐시(tempdir)를 자동 탐색."""
    if not path:
        tmp = tempfile.gettempdir()
        for cand in (os.path.join(tmp, 'learned_addr_map.json'),
                     os.path.join(tmp, 'cert_cache.db')):
            if os.path.exists(cand):
                path = cand
                print(f'학습맵 자동 탐색: {path}')
                break
    if not path:
        return {}, ''
    if not os.path.exists(path):
        sys.exit(f'학습 맵 파일 없음: {path}')
    if path.lower().endswith('.json'):
        with open(path, encoding='utf-8') as f:
            raw = json.load(f)
        learned = logic['_normalize_learned_map'](raw)
        dropped = len(raw) - len(learned)
        if dropped:
            print(f'경고: 학습맵 {len(raw)}개 중 {dropped}개 항목이 유효한 팀이 아니라 제외됨')
        if not learned:
            sys.exit(f'학습맵에 쓸 수 있는 항목이 없습니다: {path}\n'
                     '  구 형식({"access":…,"team":…}) 캐시일 수 있습니다. '
                     '파일을 지우고 cert DB(-l …/cert_cache.db)로 직접 학습시키세요.')
        return learned, path
    learned = logic['_learn_addr_map_from_cert_db'](db_path=path)
    if not learned:
        sys.exit(f'cert DB에서 학습 실패(테이블/컬럼 확인): {path}')
    return learned, path


def load_legal_dong(with_ri):
    """법정동코드 = 시도(2) + 시군구(3) + 읍면동(3) + 리(2).

    시도/시군구 명칭은 토큰 개수로 자르지 않고 상위 코드 행에서 역산한다
    (세종특별자치시처럼 시군구가 없거나 '성남시 분당구'처럼 2토큰인 경우 대응).
    """
    index = {}
    raw = []
    with open(LEGAL_DONG_TSV, encoding='utf-8') as f:
        next(f)
        for line in f:
            p = line.rstrip('\n').split('\t')
            if len(p) < 3 or p[2] != '존재':
                continue
            code, name = p[0].strip(), p[1].strip()
            if len(code) != 10:
                continue
            index[code] = name
            raw.append((code, name))

    rows = []
    for code, name in raw:
        if code[2:] == '00000000':
            continue
        if code[5:] == '00000':
            continue
        level = '읍면동' if code[8:] == '00' else '리'
        if level == '리' and not with_ri:
            continue
        시군구_full = index.get(code[:5] + '00000', '')
        # 세종특별자치시처럼 시도 단독 행이 없는 경우 첫 토큰으로 보정
        시도 = index.get(code[:2] + '00000000', '') or (시군구_full or name).split()[0]
        시군구 = 시군구_full[len(시도):].strip() if 시군구_full.startswith(시도) else 시군구_full
        prefix = 시군구_full or 시도
        하위 = name[len(prefix):].strip() if name.startswith(prefix) else name
        rows.append((code, name, level, 시도, 시군구, 하위))
    return rows


def load_skt_ops(path):
    """품질개선팀 → SKT Access운용팀."""
    src = path or SKT_OPS_TSV
    out = {}
    if not os.path.exists(src):
        return out
    with open(src, encoding='utf-8') as f:
        next(f)
        for line in f:
            p = line.rstrip('\n').split('\t')
            if len(p) >= 3:
                out[p[0]] = p[2]
    return out


def load_admin_dong(path):
    """행정동↔법정동 매핑 로드 → [(행정동코드, 시도, 시군구, 행정동명, 법정동코드), …].

    path 가 .xlsx 면 행안부 KIKmix 원본으로 보고 말소분을 걸러 읽는다.
    """
    if path and path.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        out = []
        for h, sido, sgg, hdong, bcode, _dongri, _c, malso in wb.active.iter_rows(
                min_row=2, values_only=True):
            if not hdong or malso:
                continue
            out.append((str(h), sido or '', sgg or '', hdong, str(bcode)))
        return out

    src = path or ADMIN_DONG_TSV_GZ
    if not os.path.exists(src):
        return []
    opener = gzip.open if src.endswith('.gz') else open
    out = []
    with opener(src, 'rt', encoding='utf-8') as f:
        next(f)
        for line in f:
            p = line.rstrip('\n').split('\t')
            if len(p) >= 5:
                out.append(tuple(p[:5]))
    return out


def build(args):
    api_dir = resolve_paths(args.api_dir)
    logic = load_logic()
    learned_raw, learned_src = load_learned(args.learned, logic)
    hdqt_from_addr = logic['_hdqt_from_addr']
    normalize_addr = logic['_normalize_addr']
    SEOUL_SORTED = logic['_SEOUL_GU_SORTED']
    ORG = logic['INSP_ORG_MAP']
    TEAM_TO_HDQT = logic['INSP_TEAM_TO_HDQT']
    ACCESS_TO_SKT = logic['_ACCESS_TO_SKT_HDQT']
    SKT_OPS = load_skt_ops(args.skt_ops)
    missing_ops = sorted(t for t in TEAM_TO_HDQT if t not in SKT_OPS)
    if SKT_OPS and missing_ops:
        print(f'경고: SKT 운용팀 매핑 없는 품질개선팀 {len(missing_ops)}개 — {missing_ops}')

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='2F5597')

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal='center', vertical='center')
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        return ws

    # 1) 조직도
    org_rows = []
    for access, teams in ORG.items():
        for t in teams:
            org_rows.append([ACCESS_TO_SKT.get(access, ''), access, t])
    sheet('조직도', ['SKT본부', 'access담당(본부)', '품질개선팀'], org_rows, [12, 16, 20])

    # 2) 서울 자치구 확정 규칙
    seoul_rows = [[gu, hd, tm] for gu, (hd, tm) in logic['_SEOUL_GU_TO_TEAM'].items()]
    seoul_rows.sort(key=lambda r: (r[1], r[2], r[0]))
    sheet('서울_자치구_규칙', ['자치구', 'access담당(본부)', '품질개선팀'],
          seoul_rows, [12, 16, 20])

    # 3) 법정동 단위 매핑 (메인)
    dong_rows = []
    stat_matched = 0
    for code, name, level, 시도, 시군구, 하위 in load_legal_dong(args.with_ri):
        rec = RecordingMap(learned_raw)
        hdqt, team = hdqt_from_addr(name, learned_map=rec)
        if team:
            stat_matched += 1
            if rec.hit:
                근거 = f'학습맵: {rec.hit}'
            else:
                norm = normalize_addr(name)
                gu = next((k for k, _ in SEOUL_SORTED if k in norm), '')
                근거 = f'서울 자치구 규칙: {gu}' if gu else '확정규칙'
        else:
            근거 = '미매핑'
        dong_rows.append([code, 시도, 시군구, 하위, level, name,
                          ACCESS_TO_SKT.get(hdqt, ''), hdqt or '', team or '', 근거])
    total = len(dong_rows)

    # 시군구 합의에서 벗어난 소수 판정 표시.
    # 복합키가 소수 표본(학습 임계 3건)으로 지역 합의를 뒤집는 경우가 있어
    # (예: 지하철 역사가 몰린 동 → 지하철품질개선팀) 재생성 때마다 눈에 띄게 한다.
    from collections import Counter, defaultdict
    sgg_teams = defaultdict(Counter)
    for r in dong_rows:
        if r[8]:
            sgg_teams[(r[1], r[2])][r[8]] += 1
    review = 0
    for r in dong_rows:
        counter = sgg_teams.get((r[1], r[2]))
        if not counter or not r[8]:
            r += ['', '']
            continue
        top_team, top_n = counter.most_common(1)[0]
        n, tot = counter[r[8]], sum(counter.values())
        if r[8] != top_team:
            review += 1
            r += [top_team, f'소수 {n}/{tot}']
        else:
            r += [top_team, '']

    sheet('법정동_팀매핑',
          ['법정동코드', '시도', '시군구', '읍면동', '단위', '전체주소',
           'SKT본부', 'access담당(본부)', '품질개선팀', '판정근거',
           '시군구_대표팀', '검토필요'],
          dong_rows, [13, 14, 18, 14, 8, 42, 10, 14, 20, 26, 20, 12])

    # 4) 시군구 단위 요약 (한 시군구가 여러 팀으로 갈리는지 확인용)
    agg = defaultdict(lambda: defaultdict(int))
    for r in dong_rows:
        agg[(r[1], r[2])][r[8] or '(미매핑)'] += 1
    sgg_rows = []
    for (시도, 시군구), counter in sorted(agg.items()):
        items = sorted(counter.items(), key=lambda x: -x[1])
        sgg_rows.append([시도, 시군구, len(items),
                         ', '.join(f'{t}({n})' for t, n in items),
                         items[0][0] if items else ''])
    sheet('시군구_요약', ['시도', '시군구', '팀종류수', '팀별 읍면동수', '대표팀'],
          sgg_rows, [14, 20, 10, 60, 20])

    # 4-2) 행정동 단위 매핑 — 기존 데이터가 행정동 기준일 때 조인용.
    # 행정동 하나가 여러 법정동을 관할하므로 구성 법정동의 판정을 다수결로 집계한다.
    admin_rows = []
    adm_review = 0
    adm_src = load_admin_dong(args.admin_dong)
    if adm_src:
        bjd_team = {r[0]: r[8] for r in dong_rows if r[8]}
        sgg_top = {k: c.most_common(1)[0][0] for k, c in sgg_teams.items() if c}
        adm_agg = defaultdict(Counter)
        adm_info = {}
        for hcode, sido, sgg, hdong, bcode in adm_src:
            adm_info[hcode] = (sido, sgg, hdong)
            team = bjd_team.get(bcode) or bjd_team.get(bcode[:8] + '00')
            if team:
                adm_agg[hcode][team] += 1
        for hcode, (sido, sgg, hdong) in sorted(adm_info.items()):
            counter = adm_agg.get(hcode)
            if counter:
                team, n = counter.most_common(1)[0]
                tot = sum(counter.values())
                if len(counter) == 1:
                    근거, 검토 = f'구성 법정동 {tot}개 전부 동일', ''
                else:
                    근거 = f'다수결 {n}/{tot}'
                    검토 = '검토'
                    adm_review += 1
            else:
                # 출장소(관할이 시군구 코드) 또는 신설 법정동 — 시군구 대표팀으로 폴백
                team = sgg_top.get((sido, sgg), '')
                tot = 0
                근거 = '시군구 대표팀 폴백' if team else '미매핑'
                검토 = '검토' if team else ''
                if team:
                    adm_review += 1
            ons = TEAM_TO_HDQT.get(team, '')
            admin_rows.append([hcode, sido, sgg, hdong, tot, team, ons,
                               ACCESS_TO_SKT.get(ons, ''), SKT_OPS.get(team, ''),
                               근거, 검토])
        sheet('행정동_팀매핑',
              ['행정동코드', '시도', '시군구', '행정동명', '구성법정동수',
               '품질개선팀', 'access담당(ONS)', 'access담당(SKT)', 'access운용팀(SKT)',
               '판정근거', '검토필요'],
              admin_rows, [13, 14, 18, 20, 12, 20, 16, 16, 18, 24, 10])

    # 5) 학습된 키워드 → 팀
    lrn_rows = sorted(
        ([kw, t, ons, ACCESS_TO_SKT.get(ons, ''), SKT_OPS.get(t, '')]
         for kw, t, ons in ((k, v, TEAM_TO_HDQT.get(v, ''))
                            for k, v in learned_raw.items())),
        key=lambda r: (r[3], r[2], r[1], r[0]))
    sheet('학습_키워드_팀',
          ['주소 키워드', '품질개선팀', 'access담당(ONS)', 'access담당(SKT)',
           'access운용팀(SKT)'],
          lrn_rows, [26, 20, 16, 16, 18])

    # 6) 폐지팀 치환 / 주소 약어 정규화
    sheet('폐지팀_치환', ['구(폐지) 팀명', '현행 팀명'],
          sorted(logic['_DEPRECATED_TEAM_MAP'].items()), [22, 22])
    sheet('주소약어_정규화', ['축약 표기', '정규화 표기'],
          [[a.strip(), f.strip()] for a, f in logic['_ADDR_ABBR_MAP'].items()], [14, 20])

    # 0) 안내 시트
    ws = wb['Sheet']
    ws.title = '안내'
    guide = [
        ['주소 → 품질개선팀 매핑 로직'],
        [''],
        ['생성 스크립트', 'scripts/export_addr_team_map.py'],
        ['로직 원본', f'{api_dir}/routers/inspection.py  _hdqt_from_addr()'],
        ['법정동 원본', LEGAL_DONG_TSV],
        ['학습맵 입력', learned_src or '(없음 — 서울 확정규칙만 반영됨)'],
        [''],
        ['■ 판정 순서'],
        ['1', '주소 정규화 — 앞쪽 괄호 제거, "서울"→"서울특별시" 등 축약 표기 확장 (시트: 주소약어_정규화)'],
        ['2', '서울이면 자치구 확정 규칙표로 즉시 결정 (시트: 서울_자치구_규칙). 가장 긴 키워드 우선'],
        ['3', '서울이 아니면 주소에서 [시/군/구/읍/면/동] 토큰을 뽑아 후보 키를 만든다'],
        ['',  '   복합키: "시 구", "시 군", "시 동", "군 읍", "군 면", "구 동"'],
        ['',  '   단일키: 시/군/구/읍/면 (단독 "동"은 동명이 전국에 중복되어 제외)'],
        ['4', '후보키를 길이 내림차순으로 학습맵에 조회 → 첫 히트의 팀으로 확정 (시트: 학습_키워드_팀)'],
        ['5', '팀이 정해지면 조직도로 본부를 역산 (시트: 조직도)'],
        ['6', '끝까지 못 찾으면 기존 본부값 유지, 팀은 공란'],
        [''],
        ['■ 학습맵이란'],
        ['',  'cert DB(무선국 인증 데이터)의 주소(zpwiadr) + 실제 담당팀(ons_team_nm) 쌍을 집계해'],
        ['',  '키워드별 최빈 팀을 채택한 표. 동일 키워드 3건 미만은 노이즈로 버린다.'],
        ['',  '즉 읍면동 단위 매핑은 하드코딩이 아니라 운영 데이터에서 학습된 결과다.'],
        [''],
        ['■ 검토필요 컬럼'],
        ['',  '같은 시군구 안에서 대표팀과 다른 팀으로 판정된 행에 "소수 n/전체" 표시.'],
        ['',  '실제로 팀 경계가 시군구를 가르는 정상 케이스도 있지만(예: 인천 중구 영종도),'],
        ['',  '"소수 1/N" 처럼 극단적으로 적은 건은 소수 표본이 지역 합의를 뒤집은 오매핑일 수 있다.'],
        ['',  '학습 임계가 3건이라 특정 동에 지하철 역사 등이 몰리면 발생한다. 눈으로 확인할 것.'],
        [''],
        ['■ 행정동_팀매핑 시트'],
        ['',  '기존 데이터가 행정동 기준일 때 조인용. 행정동↔법정동은 1:1이 아니라 다대다다.'],
        ['',  '(청운효자동 ⊃ 청운동·신교동·궁정동…  /  신림동 → 여러 행정동으로 분할)'],
        ['',  '그래서 행정동이 관할하는 법정동들의 판정을 모아 다수결로 팀을 정한다.'],
        ['',  '매핑 원본: 행정안전부 행정기관(행정동) 및 관할구역(법정동) KIKmix, 2026.3.25 시행'],
        ['',  '출장소나 신설 법정동처럼 구성 법정동이 안 잡히는 건은 시군구 대표팀으로 폴백하고'],
        ['',  '판정근거에 표시한다. 검토필요 표시된 행만 확인하면 된다.'],
        [''],
        ['■ 커버리지'],
        ['대상 읍면동 수', total],
        ['팀 확정', stat_matched],
        ['미매핑', total - stat_matched],
        ['검토필요(소수 판정)', review],
        ['학습 키워드 수', len(learned_raw)],
        ['행정동 수', len(admin_rows)],
        ['행정동 검토필요', adm_review],
    ]
    for row in guide:
        ws.append(row)
    ws['A1'].font = Font(bold=True, size=14)
    for r in ws.iter_rows(min_col=1, max_col=1):
        if r[0].value and str(r[0].value).startswith('■'):
            r[0].font = Font(bold=True, size=11)
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 100

    out = args.out or os.path.join(ROOT, 'docs', '주소-팀_매핑.xlsx')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f'생성 완료: {out}')
    print(f'  읍면동 {total}건 중 팀 확정 {stat_matched}건 / 미매핑 {total - stat_matched}건')
    print(f'  학습 키워드 {len(learned_raw)}개')
    print(f'  검토필요(시군구 대표팀과 다른 소수 판정) {review}건')
    if admin_rows:
        print(f'  행정동 {len(admin_rows)}건 / 검토필요 {adm_review}건')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('-o', '--out', help='출력 xlsx 경로')
    ap.add_argument('-l', '--learned',
                    help='학습맵 (.json 캐시 또는 cert_cache.db). 생략 시 임시디렉터리 자동 탐색')
    ap.add_argument('--api-dir', help='inspection.py / legal_dong_code.tsv 가 있는 디렉터리')
    ap.add_argument('--skt-ops',
                    help='품질개선팀→SKT Access운용팀 tsv. '
                         '생략 시 scripts/data/skt_ops_team_map.tsv 사용')
    ap.add_argument('--admin-dong',
                    help='행정동↔법정동 매핑. 행안부 KIKmix xlsx 또는 tsv(.gz). '
                         '생략 시 scripts/data/admin_dong_map.tsv.gz 사용')
    ap.add_argument('--with-ri', action='store_true', help='리(里) 단위까지 포함')
    build(ap.parse_args())
