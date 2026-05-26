"""
callname - 호출명칭 DB 관리 엔드포인트

담당 도메인: 호출명칭 DB 업로드/매칭/샘플 관리
주요 의존성: core.auth, core.config, core.s3, core.cert_cache, core.utils
엔드포인트:
    POST /callname/upload-csv
    GET  /callname/upload-job/{job_id}
    GET  /callname/db-status
    GET  /callname/db-preview
    POST /callname/upload-raw
    POST /callname/upload-complete
    GET  /callname/upload/{upload_id}/analysis
    POST /callname/upload
    POST /callname/upload/{upload_id}/column-values
    POST /callname/upload/{upload_id}/preview
    POST /callname/process
    GET  /callname/process/{process_id}/stream
    GET  /callname/process/{process_id}/download
    GET  /callname/sample-template
    POST /callname/sample-template
    GET  /callname/sample-template/download
    DELETE /callname/sample-template
"""

import asyncio
import logging
import os
import re
import tempfile as _tempfile
import time as _time_mod
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from typing import Dict, Optional
from urllib.parse import quote

from fastapi import APIRouter, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse

from core.auth import _verify_auth, _require_role, _record_audit_log_sync
from core.cert_cache import (
    _cert_cache_force_rebuild, _query_callname_db,
)
from core.config import (
    S3_BUCKET_NAME, MAX_DS_UPLOAD_SIZE,
    CALLNAME_CSV_PREFIX, CALLNAME_USE_COLS,
    CALLNAME_SESSION_TTL, CALLNAME_MAX_SESSIONS,
)
from core.s3 import get_s3_client
from core.utils import _check_memory, _log_mem, _release_memory, _check_rate_limit, _get_pandas

router = APIRouter(tags=["callname"])
logger = logging.getLogger(__name__)

# ── 선택적 의존성 ──────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.utils import get_column_letter, column_index_from_string
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    get_column_letter = None
    column_index_from_string = None

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

# ── 호출명칭 설정 상수 ──────────────────────────────────────────
CALLNAME_POSSIBLE_CALLNAME_COLS = ["호출명칭", "callname", "CALLNAME", "호출명", "call_name"]
CALLNAME_POSSIBLE_TONGSI_COLS = ["통시", "통합시설코드", "zpcode"]
CALLNAME_POSSIBLE_ZPWINA_COLS = ["zpwina", "ZPWINA", "Zpwina", "호출명칭", "호출명"]
CALLNAME_POSSIBLE_ZPWINO_COLS = ["zpwino", "ZPWINO", "Zpwino", "허가번호", "허가번호O"]
CALLNAME_POSSIBLE_ACCESS_COLS = ["Access담당", "access담당", "ACCESS담당"]
CALLNAME_POSSIBLE_QUALITY_COLS = ["품질개선팀", "품질개선", "QI팀"]
CALLNAME_DB_TO_EXCEL_MAP = {
    "area_hdofc_nm": CALLNAME_POSSIBLE_ACCESS_COLS,
    "ons_team_nm": CALLNAME_POSSIBLE_QUALITY_COLS,
    "zpcode": CALLNAME_POSSIBLE_TONGSI_COLS,
}
CALLNAME_SAMPLE_PREFIX = "callname-sample/"
CALLNAME_SAMPLE_MAX_SIZE = 20 * 1024 * 1024  # 20MB
CALLNAME_SAMPLE_ALLOWED_EXTS = ("xlsx", "xls")

# ── 상태 저장 ─────────────────────────────────────────────────
_callname_db_row_count = 0
_callname_sessions: Dict[str, dict] = {}
_callname_upload_jobs: Dict[str, dict] = {}
_bounded_executor = ThreadPoolExecutor(max_workers=2)

# ── 통시 NA 패턴 ──────────────────────────────────────────────
_TONGSI_NA_VALUES = frozenset({
    "", "#n/a", "#na", "n/a", "na", "nan", "#ref!", "#value!", "#null!",
    "null", "none", "-", "--",
})


def _is_tongsi_empty(val: str) -> bool:
    stripped = val.strip()
    if not stripped:
        return True
    return stripped.lower() in _TONGSI_NA_VALUES


# ── xlsx 헬퍼 ─────────────────────────────────────────────────

def _col_to_idx(col_letter: str) -> int:
    r = 0
    for c in col_letter:
        r = r * 26 + (ord(c) - 64)
    return r - 1


def _resolve_xlsx_sheet_path(zf, sheet_name: str) -> Optional[str]:
    import xml.etree.ElementTree as ET
    try:
        wb_xml = zf.read("xl/workbook.xml")
        wb_root = ET.fromstring(wb_xml)
        r_id = None
        for el in wb_root.iter():
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "sheet" and el.get("name") == sheet_name:
                for attr_key in el.attrib:
                    if attr_key.endswith("}id") or attr_key == "r:id":
                        r_id = el.attrib[attr_key]
                        break
                break
        del wb_xml, wb_root
        if not r_id:
            return None
        rels_xml = zf.read("xl/_rels/workbook.xml.rels")
        rels_root = ET.fromstring(rels_xml)
        for rel in rels_root:
            if rel.get("Id") == r_id:
                target = rel.get("Target", "")
                del rels_xml, rels_root
                if target.startswith("/"):
                    return target[1:]
                return f"xl/{target}"
        del rels_xml, rels_root
    except Exception:
        pass
    return None


def _iter_xlsx_rows_light(xlsx_path: str, sheet_name: str = None, *,
                          ss_cache_path: str = None, ss_offsets_bytes: bytes = None):
    """xlsx → (0-based_row_num, [str, ...]) 스트리밍 제너레이터."""
    import xml.etree.ElementTree as ET
    import struct
    import mmap as _mmap_mod
    from array import array

    _owns_ss = ss_cache_path is None
    ss_tmp_path = ss_cache_path
    ss_mmap_obj = None
    ss_fh = None

    try:
        with zipfile.ZipFile(xlsx_path, "r") as zf:
            if ss_cache_path and ss_offsets_bytes:
                ss_offsets = array("Q")
                ss_offsets.frombytes(ss_offsets_bytes)
            else:
                ss_offsets = array("Q")
                ss_names = [n for n in zf.namelist() if n.endswith("sharedStrings.xml")]
                if ss_names:
                    ss_tmp = _tempfile.NamedTemporaryFile(delete=False, suffix=".ss")
                    ss_tmp_path = ss_tmp.name
                    with zf.open(ss_names[0]) as ssf:
                        for _, elem in ET.iterparse(ssf, events=("end",)):
                            tag = elem.tag.rsplit("}", 1)[-1]
                            if tag == "si":
                                parts = []
                                for ch in elem.iter():
                                    if ch.tag.rsplit("}", 1)[-1] == "t" and ch.text:
                                        parts.append(ch.text)
                                text = "".join(parts)
                                encoded = text.encode("utf-8")
                                ss_offsets.append(ss_tmp.tell())
                                ss_tmp.write(struct.pack("<I", len(encoded)))
                                ss_tmp.write(encoded)
                                elem.clear()
                    ss_tmp.close()

            if ss_tmp_path:
                file_size = os.path.getsize(ss_tmp_path)
                if file_size > 0:
                    ss_fh = open(ss_tmp_path, "rb")
                    ss_mmap_obj = _mmap_mod.mmap(ss_fh.fileno(), 0, access=_mmap_mod.ACCESS_READ)

            def _get_ss(idx):
                if ss_mmap_obj is not None and 0 <= idx < len(ss_offsets):
                    import struct as _s
                    offset = ss_offsets[idx]
                    length = _s.unpack_from("<I", ss_mmap_obj, offset)[0]
                    start = offset + 4
                    return ss_mmap_obj[start:start + length].decode("utf-8")
                return ""

            if sheet_name:
                sp = _resolve_xlsx_sheet_path(zf, sheet_name)
                if not sp:
                    return
            else:
                sheets = sorted([n for n in zf.namelist() if "worksheets/sheet" in n])
                sp = sheets[0] if sheets else "xl/worksheets/sheet1.xml"

            _cr = re.compile(r"([A-Z]+)")
            cells = []

            with zf.open(sp) as sf:
                for _, elem in ET.iterparse(sf, events=("end",)):
                    tag = elem.tag.rsplit("}", 1)[-1]

                    if tag == "c":
                        ct = elem.get("t", "")
                        val = ""
                        if ct == "s":
                            for ch in elem:
                                if ch.tag.rsplit("}", 1)[-1] == "v" and ch.text:
                                    si = int(ch.text)
                                    val = _get_ss(si)
                                    break
                        elif ct == "inlineStr":
                            for ch in elem.iter():
                                if ch.tag.rsplit("}", 1)[-1] == "t" and ch.text:
                                    val = ch.text
                                    break
                        else:
                            for ch in elem:
                                if ch.tag.rsplit("}", 1)[-1] == "v":
                                    val = ch.text or ""
                                    break
                        r_attr = elem.get("r", "")
                        m = _cr.match(r_attr)
                        if m:
                            ci = _col_to_idx(m.group(1))
                            while len(cells) <= ci:
                                cells.append("")
                            cells[ci] = val
                        elem.clear()

                    elif tag == "row":
                        rn = int(elem.get("r", "0")) - 1
                        yield (rn, cells)
                        cells = []
                        elem.clear()

    finally:
        if ss_mmap_obj is not None:
            ss_mmap_obj.close()
        if ss_fh is not None:
            ss_fh.close()
        if _owns_ss and ss_tmp_path:
            try:
                os.remove(ss_tmp_path)
            except Exception:
                pass


def _parse_xlsx_header_fast(xlsx_path: str) -> dict:
    """xlsx 헤더(첫 행) + 행 수만 초고속 추출."""
    import xml.etree.ElementTree as ET
    _cr_col = re.compile(r"([A-Z]+)")

    with zipfile.ZipFile(xlsx_path, "r") as zf:
        sheets = sorted([n for n in zf.namelist() if "worksheets/sheet" in n])
        sp = sheets[0] if sheets else "xl/worksheets/sheet1.xml"

        total_rows = 0
        header_cells = []

        with zf.open(sp) as sf:
            for _, elem in ET.iterparse(sf, events=("end",)):
                tag = elem.tag.rsplit("}", 1)[-1]

                if tag == "dimension":
                    ref = elem.get("ref", "")
                    if ":" in ref:
                        m = re.search(r"(\d+)$", ref.split(":")[1])
                        if m:
                            total_rows = int(m.group(1)) - 1
                    elem.clear()
                    continue

                if tag == "c":
                    r_attr = elem.get("r", "")
                    if r_attr and r_attr[-1] == "1" and re.match(r"^[A-Z]+1$", r_attr):
                        ct = elem.get("t", "")
                        m = _cr_col.match(r_attr)
                        ci = _col_to_idx(m.group(1)) if m else -1

                        if ct == "s":
                            for ch in elem:
                                if ch.tag.rsplit("}", 1)[-1] == "v" and ch.text:
                                    header_cells.append((ci, "s", int(ch.text)))
                                    break
                        elif ct == "inlineStr":
                            for ch in elem.iter():
                                if ch.tag.rsplit("}", 1)[-1] == "t" and ch.text:
                                    header_cells.append((ci, "v", ch.text))
                                    break
                        else:
                            for ch in elem:
                                if ch.tag.rsplit("}", 1)[-1] == "v":
                                    header_cells.append((ci, "v", ch.text or ""))
                                    break
                    elem.clear()
                    continue

                if tag == "row":
                    rn = int(elem.get("r", "0"))
                    elem.clear()
                    if rn >= 2:
                        break
                elif tag not in ("v", "t"):
                    elem.clear()

        needed = {val for _, typ, val in header_cells if typ == "s"}
        ss_map = {}
        if needed:
            max_idx = max(needed)
            ss_names = [n for n in zf.namelist() if n.endswith("sharedStrings.xml")]
            if ss_names:
                idx = 0
                with zf.open(ss_names[0]) as ssf:
                    for _, elem in ET.iterparse(ssf, events=("end",)):
                        tag = elem.tag.rsplit("}", 1)[-1]
                        if tag == "si":
                            if idx in needed:
                                parts = []
                                for ch in elem.iter():
                                    if ch.tag.rsplit("}", 1)[-1] == "t" and ch.text:
                                        parts.append(ch.text)
                                ss_map[idx] = "".join(parts)
                            elem.clear()
                            idx += 1
                            if idx > max_idx:
                                break

        if header_cells:
            max_ci = max(ci for ci, _, _ in header_cells)
            columns = [""] * (max_ci + 1)
            for ci, typ, val in header_cells:
                columns[ci] = ss_map.get(val, "") if typ == "s" else str(val)
        else:
            columns = []

        while columns and not columns[-1].strip():
            columns.pop()

        return {"columns": columns, "total_rows": total_rows}


def _detect_column(df_columns, candidates):
    """컬럼 목록에서 후보 이름과 일치하는 첫 번째 컬럼명 반환. 3단계 매칭."""
    col_list = list(df_columns)
    for name in candidates:
        if name in col_list:
            return name
    stripped_map = {c.strip().lower(): c for c in col_list if c.strip()}
    for name in candidates:
        key = name.strip().lower()
        if key in stripped_map:
            return stripped_map[key]
    for name in candidates:
        nl = name.strip().lower()
        if not nl:
            continue
        for c in col_list:
            if nl in c.strip().lower():
                return c
    return None


def _s3_to_tempfile(s3_key: str, suffix: str = ".tmp") -> str:
    """S3 파일을 디스크 임시파일로 스트리밍 다운로드."""
    obj = get_s3_client().get_object(Bucket=S3_BUCKET_NAME, Key=s3_key)
    tmp = _tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    body = obj["Body"]
    try:
        while True:
            chunk = body.read(1024 * 1024)
            if not chunk:
                break
            tmp.write(chunk)
    finally:
        body.close()
    tmp.close()
    return tmp.name


def _cleanup_callname_session_files(sess: dict):
    """세션의 캐시/임시 파일 정리 (디스크 + S3) + 대용량 데이터 해제."""
    s3_temp = sess.get("s3_temp_key")
    if s3_temp:
        try:
            get_s3_client().delete_object(Bucket=S3_BUCKET_NAME, Key=s3_temp)
        except Exception:
            pass
    for path_key in ("cached_xlsx_path", "cached_ss_path"):
        p = sess.get(path_key)
        if p:
            try:
                os.remove(p)
            except Exception:
                pass
    for data_key in ("filter_cache_rows", "filter_cache_row_indices",
                      "column_stats", "cached_ss_offsets",
                      "original_row_indices", "row_zpwina_list", "row_zpwino_list",
                      "zpwina_values", "zpwino_values"):
        sess.pop(data_key, None)


def _cleanup_callname_sessions():
    """만료된 세션 + S3/디스크 임시 파일 정리."""
    now = _time_mod.time()
    expired = [k for k, v in _callname_sessions.items()
               if (now - v.get("created_at_ts", 0)) > CALLNAME_SESSION_TTL]
    for k in expired:
        _cleanup_callname_session_files(_callname_sessions[k])
        del _callname_sessions[k]


def _analyze_callname_bg(upload_id: str):
    """Background: xlsx 단일 ZIP 오픈 → SS캐시 빌드 + 전행 스캔 통합."""
    import xml.etree.ElementTree as ET
    import struct
    import mmap as _mmap_mod
    from array import array
    from collections import Counter

    sess = _callname_sessions.get(upload_id)
    if not sess or sess.get("status") != "uploaded":
        return
    try:
        sess["analysis_status"] = "processing"

        tmp_path = sess.get("cached_xlsx_path")
        if not tmp_path or not os.path.exists(tmp_path):
            tmp_path = _s3_to_tempfile(sess["s3_temp_key"], f".{sess['ext']}")
            sess["cached_xlsx_path"] = tmp_path

        columns = []
        filtered_rows = 0
        total_rows = 0
        callname_set = set()
        col_counters = []
        filter_cache_rows = []
        filter_cache_row_indices = []
        ss_cache_path = None
        ss_offsets_bytes = None

        if sess.get("ext") == "xlsx":
            _cr = re.compile(r"([A-Z]+)")
            ss_offsets = array("Q")
            ss_tmp_path = None
            ss_mmap_obj = None
            ss_fh = None

            try:
                with zipfile.ZipFile(tmp_path, "r") as zf:
                    ss_names = [n for n in zf.namelist() if n.endswith("sharedStrings.xml")]
                    if ss_names:
                        ss_tmp = _tempfile.NamedTemporaryFile(delete=False, suffix=".ss")
                        ss_tmp_path = ss_tmp.name
                        with zf.open(ss_names[0]) as ssf:
                            for _, elem in ET.iterparse(ssf, events=("end",)):
                                tag = elem.tag.rsplit("}", 1)[-1]
                                if tag == "si":
                                    parts = []
                                    for ch in elem.iter():
                                        if ch.tag.rsplit("}", 1)[-1] == "t" and ch.text:
                                            parts.append(ch.text)
                                    text = "".join(parts)
                                    encoded = text.encode("utf-8")
                                    ss_offsets.append(ss_tmp.tell())
                                    ss_tmp.write(struct.pack("<I", len(encoded)))
                                    ss_tmp.write(encoded)
                                    elem.clear()
                        ss_tmp.close()

                        file_size = os.path.getsize(ss_tmp_path)
                        if file_size > 0:
                            ss_fh = open(ss_tmp_path, "rb")
                            ss_mmap_obj = _mmap_mod.mmap(
                                ss_fh.fileno(), 0, access=_mmap_mod.ACCESS_READ)

                    def _get_ss(idx):
                        if ss_mmap_obj is not None and 0 <= idx < len(ss_offsets):
                            offset = ss_offsets[idx]
                            length = struct.unpack_from("<I", ss_mmap_obj, offset)[0]
                            start = offset + 4
                            return ss_mmap_obj[start:start + length].decode("utf-8")
                        return ""

                    sheets = sorted([n for n in zf.namelist()
                                     if "worksheets/sheet" in n])
                    sp = sheets[0] if sheets else "xl/worksheets/sheet1.xml"
                    cells_buf = []
                    row_num = 0
                    tongsi_idx = -1
                    callname_idx = -1
                    zpwina_idx = -1
                    zpwino_idx = -1

                    with zf.open(sp) as sf:
                        for _, elem in ET.iterparse(sf, events=("end",)):
                            tag = elem.tag.rsplit("}", 1)[-1]

                            if tag == "c":
                                ct = elem.get("t", "")
                                val = ""
                                if ct == "s":
                                    for ch in elem:
                                        if ch.tag.rsplit("}", 1)[-1] == "v" and ch.text:
                                            val = _get_ss(int(ch.text))
                                            break
                                elif ct == "inlineStr":
                                    for ch in elem.iter():
                                        if ch.tag.rsplit("}", 1)[-1] == "t" and ch.text:
                                            val = ch.text
                                            break
                                else:
                                    for ch in elem:
                                        if ch.tag.rsplit("}", 1)[-1] == "v":
                                            val = ch.text or ""
                                            break
                                r_attr = elem.get("r", "")
                                m = _cr.match(r_attr)
                                if m:
                                    ci = _col_to_idx(m.group(1))
                                    while len(cells_buf) <= ci:
                                        cells_buf.append("")
                                    cells_buf[ci] = val
                                elem.clear()

                            elif tag == "row":
                                rn = int(elem.get("r", "0")) - 1
                                if rn == 0:
                                    columns = list(cells_buf)
                                    while columns and not columns[-1].strip():
                                        columns.pop()
                                    tongsi_col = _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS)
                                    callname_col = _detect_column(columns, CALLNAME_POSSIBLE_CALLNAME_COLS)
                                    zpwina_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINA_COLS)
                                    zpwino_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINO_COLS)
                                    tongsi_idx = columns.index(tongsi_col) if tongsi_col and tongsi_col in columns else -1
                                    callname_idx = columns.index(callname_col) if callname_col and callname_col in columns else -1
                                    zpwina_idx = columns.index(zpwina_col) if zpwina_col and zpwina_col in columns else -1
                                    zpwino_idx = columns.index(zpwino_col) if zpwino_col and zpwino_col in columns else -1
                                    col_counters = [Counter() for _ in columns]
                                else:
                                    total_rows += 1
                                    tongsi_val = cells_buf[tongsi_idx] if tongsi_idx >= 0 and tongsi_idx < len(cells_buf) else ""
                                    if _is_tongsi_empty(tongsi_val):
                                        filtered_rows += 1
                                        filter_cache_rows.append(list(cells_buf) + [""] * max(0, len(columns) - len(cells_buf)))
                                        filter_cache_row_indices.append(rn + 1)
                                        cn = cells_buf[callname_idx].strip() if callname_idx >= 0 and callname_idx < len(cells_buf) else ""
                                        if cn:
                                            callname_set.add(cn)
                                    for ci, counter in enumerate(col_counters):
                                        v = cells_buf[ci] if ci < len(cells_buf) else ""
                                        if v:
                                            counter[v] += 1
                                cells_buf = []
                                elem.clear()

                ss_cache_path = ss_tmp_path
                ss_offsets_bytes = ss_offsets.tobytes() if ss_offsets else None

            finally:
                if ss_mmap_obj is not None:
                    ss_mmap_obj.close()
                if ss_fh is not None:
                    ss_fh.close()

        else:
            if HAS_XLRD:
                xls_book = xlrd.open_workbook(tmp_path)
                ws = xls_book.sheet_by_index(0)
                if ws.nrows > 0:
                    columns = [str(ws.cell_value(0, c)) for c in range(ws.ncols)]
                    tongsi_col = _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS)
                    callname_col = _detect_column(columns, CALLNAME_POSSIBLE_CALLNAME_COLS)
                    zpwina_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINA_COLS)
                    zpwino_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINO_COLS)
                    tongsi_idx = columns.index(tongsi_col) if tongsi_col and tongsi_col in columns else -1
                    callname_idx = columns.index(callname_col) if callname_col and callname_col in columns else -1
                    zpwina_idx = columns.index(zpwina_col) if zpwina_col and zpwina_col in columns else -1
                    zpwino_idx = columns.index(zpwino_col) if zpwino_col and zpwino_col in columns else -1
                    for r in range(1, ws.nrows):
                        total_rows += 1
                        row_vals = [str(ws.cell_value(r, c)) for c in range(ws.ncols)]
                        tongsi_val = row_vals[tongsi_idx] if tongsi_idx >= 0 else ""
                        if _is_tongsi_empty(tongsi_val):
                            filtered_rows += 1
                            filter_cache_rows.append(row_vals)
                            filter_cache_row_indices.append(r + 1)
                            cn = row_vals[callname_idx].strip() if callname_idx >= 0 and callname_idx < len(row_vals) else ""
                            if cn:
                                callname_set.add(cn)
                xls_book.release_resources()

        column_stats = {}
        for ci, counter in enumerate(col_counters):
            if ci < len(columns):
                col_name = columns[ci]
                if col_name and counter:
                    column_stats[col_name] = [{"value": v, "count": c} for v, c in counter.most_common(100)]

        sess_ref = _callname_sessions.get(upload_id)
        if sess_ref:
            sess_ref.update({
                "columns": columns,
                "total_rows": total_rows,
                "filtered_rows": filtered_rows,
                "target_callnames": len(callname_set),
                "callname_col": _detect_column(columns, CALLNAME_POSSIBLE_CALLNAME_COLS),
                "tongsi_col": _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS),
                "zpwina_col": _detect_column(columns, CALLNAME_POSSIBLE_ZPWINA_COLS),
                "zpwino_col": _detect_column(columns, CALLNAME_POSSIBLE_ZPWINO_COLS),
                "filter_cache_rows": filter_cache_rows,
                "filter_cache_row_indices": filter_cache_row_indices,
                "column_stats": column_stats,
                "cached_ss_path": ss_cache_path,
                "cached_ss_offsets": ss_offsets_bytes,
                "analysis_status": "complete",
            })

    except Exception as e:
        logger.error(f"호출명칭 백그라운드 분석 실패: {e}")
        sess_ref = _callname_sessions.get(upload_id)
        if sess_ref:
            sess_ref["analysis_status"] = "error"


def _process_callname_upload_sync(job_id: str, tmp_path: str, filename: str,
                                   ext: str, replace: bool):
    """백그라운드: 호출명칭 DB 파일 파싱 → S3 업로드."""
    import csv as _csv_mod
    global _callname_db_row_count
    job = _callname_upload_jobs[job_id]
    filtered_paths = []
    try:
        job["stage"] = "파일 분석 중..."
        job["percent"] = 10
        base_name = filename.rsplit(".", 1)[0].replace(" ", "_")

        pd = _get_pandas()

        if ext == "csv":
            filtered_path = tmp_path + ".filtered.csv"
            filtered_paths.append(("csv", filtered_path))
            first_chunk = True
            with open(filtered_path, "w", encoding="utf-8", newline="") as f:
                for chunk_df in pd.read_csv(
                    tmp_path, dtype=str, na_filter=False, chunksize=50000
                ):
                    avail = [c for c in CALLNAME_USE_COLS if c in chunk_df.columns]
                    if not avail:
                        del chunk_df
                        continue
                    chunk_df[avail].to_csv(f, index=False, header=first_chunk)
                    first_chunk = False
                    del chunk_df
            job["stage"] = "CSV 필터링 완료"
            job["percent"] = 50

        elif ext == "xlsx":
            logger.info(f"호출명칭 xlsx 파싱 시작: {filename}")
            wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            logger.info(f"호출명칭 xlsx 시트 목록: {sheet_names}")
            total_sheets = len(sheet_names)

            for sheet_idx, sn in enumerate(sheet_names):
                pct = 10 + int(40 * sheet_idx / max(total_sheets, 1))
                job["stage"] = f"시트 '{sn}' 처리 중... ({sheet_idx+1}/{total_sheets})"
                job["percent"] = pct

                ws = wb[sn]
                filtered_path = tmp_path + f".sheet{sheet_idx}.csv"
                header_row = None
                avail_indices = []
                row_count_sheet = 0
                try:
                    with open(filtered_path, "w", encoding="utf-8", newline="") as f:
                        writer = _csv_mod.writer(f)
                        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                            if row_idx == 0:
                                header_row = [str(c) if c is not None else "" for c in row]
                                avail_indices = [i for i, h in enumerate(header_row) if h in CALLNAME_USE_COLS]
                                if not avail_indices:
                                    break
                                writer.writerow([header_row[i] for i in avail_indices])
                                continue
                            writer.writerow([str(row[i]) if i < len(row) and row[i] is not None else "" for i in avail_indices])
                            row_count_sheet += 1
                            if row_count_sheet % 100000 == 0:
                                job["stage"] = f"시트 '{sn}': {row_count_sheet:,}행 처리 중..."
                                logger.info(job["stage"])
                except Exception as sheet_err:
                    logger.error(f"호출명칭 시트 '{sn}' 처리 오류: {sheet_err}")
                    try:
                        os.unlink(filtered_path)
                    except OSError:
                        pass
                    continue
                if avail_indices:
                    filtered_paths.append((f"sheet_{sn}", filtered_path))
                else:
                    try:
                        os.unlink(filtered_path)
                    except OSError:
                        pass
            wb.close()
            del wb
            job["percent"] = 50

        elif ext == "xls":
            if not HAS_XLRD:
                raise RuntimeError("xlrd 미설치")
            xls_book = xlrd.open_workbook(tmp_path)
            for sheet_idx in range(xls_book.nsheets):
                ws = xls_book.sheet_by_index(sheet_idx)
                sn = ws.name
                if ws.nrows == 0:
                    continue
                header_row = [str(ws.cell_value(0, c)) for c in range(ws.ncols)]
                avail_indices = [i for i, h in enumerate(header_row) if h in CALLNAME_USE_COLS]
                if not avail_indices:
                    continue
                filtered_path = tmp_path + f".sheet{sheet_idx}.csv"
                filtered_paths.append((f"sheet_{sn}", filtered_path))
                with open(filtered_path, "w", encoding="utf-8", newline="") as f:
                    writer = _csv_mod.writer(f)
                    writer.writerow([header_row[i] for i in avail_indices])
                    for r in range(1, ws.nrows):
                        writer.writerow([str(ws.cell_value(r, i)) for i in avail_indices])
            xls_book.release_resources()
            del xls_book
            job["percent"] = 50

        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        _release_memory()

        if not filtered_paths:
            job["status"] = "failed"
            job["stage"] = "필요한 컬럼이 포함된 시트가 없습니다."
            job["percent"] = 100
            return

        job["stage"] = "S3에 업로드 중..."
        job["percent"] = 60

        if replace:
            try:
                resp = get_s3_client().list_objects_v2(
                    Bucket=S3_BUCKET_NAME, Prefix=CALLNAME_CSV_PREFIX)
                for obj in resp.get("Contents", []):
                    get_s3_client().delete_object(
                        Bucket=S3_BUCKET_NAME, Key=obj["Key"])
                logger.info("호출명칭 DB 기존 파일 전체 삭제 (replace 모드)")
            except Exception:
                pass

        uploaded_keys = []
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        for i, (label, fpath) in enumerate(filtered_paths):
            s3_key = f"{CALLNAME_CSV_PREFIX}{base_name}_{label}_{timestamp}.csv"
            file_size = os.path.getsize(fpath)
            with open(fpath, "rb") as f:
                get_s3_client().put_object(Bucket=S3_BUCKET_NAME, Key=s3_key, Body=f)
            uploaded_keys.append(s3_key)
            pct = 60 + int(25 * (i + 1) / len(filtered_paths))
            job["stage"] = f"S3 업로드 중... ({i+1}/{len(filtered_paths)})"
            job["percent"] = pct
            logger.info(f"호출명칭 DB 업로드: {s3_key} ({file_size:,} bytes)")

        job["stage"] = "행 수 집계 중..."
        job["percent"] = 90
        uploaded_rows = 0
        for _, fpath in filtered_paths:
            with open(fpath, encoding="utf-8") as cnt_f:
                uploaded_rows += sum(1 for _ in cnt_f) - 1

        _callname_db_row_count = uploaded_rows
        import threading as _th
        _th.Thread(target=_cert_cache_force_rebuild, daemon=True).start()

        file_count = len(filtered_paths)
        job["status"] = "completed"
        job["stage"] = "완료"
        job["percent"] = 100
        job["result"] = {
            "message": f"DB 업로드 완료 ({file_count}개 파일, 총 {uploaded_rows:,}행)",
            "files": uploaded_keys,
            "total_rows": uploaded_rows,
        }
        logger.info(f"호출명칭 DB 업로드 완료: {file_count}개 파일, {uploaded_rows:,}행")

    except Exception as e:
        logger.error(f"호출명칭 DB 업로드 실패: {e}")
        job["status"] = "failed"
        job["stage"] = f"업로드 실패: {str(e)[:200]}"
        job["percent"] = 100
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        for _, fpath in filtered_paths:
            try:
                os.unlink(fpath)
            except OSError:
                pass
        _release_memory()


# ── 엔드포인트 ────────────────────────────────────────────────

@router.post("/callname/upload-csv")
async def callname_upload_csv(
    request: Request,
    file: UploadFile = File(...),
    replace: bool = Query(False, description="True면 기존 DB 전체 교체, False면 추가/병합"),
):
    """관리자: 파일 → 디스크 저장 → jobId 즉시 반환 → 백그라운드 처리."""
    await _require_role(request, {"admin"})
    _check_memory("호출명칭 DB 업로드")
    if not file.filename.lower().endswith((".csv", ".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="CSV 또는 Excel 파일만 가능합니다.")
    _get_pandas()

    ext = file.filename.rsplit(".", 1)[-1].lower()

    with _tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
        tmp_path = tmp.name
        while True:
            chunk = await file.read(8 * 1024 * 1024)
            if not chunk:
                break
            tmp.write(chunk)

    job_id = str(uuid.uuid4())
    _callname_upload_jobs[job_id] = {
        "status": "processing",
        "stage": "파일 수신 완료, 처리 시작...",
        "percent": 5,
        "filename": file.filename,
        "replace": replace,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    asyncio.get_event_loop().run_in_executor(
        _bounded_executor, _process_callname_upload_sync,
        job_id, tmp_path, file.filename, ext, replace,
    )

    return {"success": True, "jobId": job_id}


@router.get("/callname/upload-job/{job_id}")
async def callname_upload_job_status(job_id: str, request: Request):
    """호출명칭 DB 업로드 잡 상태 조회."""
    await _verify_auth(request)
    job = _callname_upload_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return {
        "status": job["status"],
        "stage": job["stage"],
        "percent": job["percent"],
        "result": job.get("result"),
    }


@router.get("/callname/db-status")
async def callname_db_status(request: Request):
    """호출명칭 DB 상태 조회 (S3 파일 목록 기반)."""
    await _verify_auth(request)
    files = []
    total_size = 0
    try:
        resp = get_s3_client().list_objects_v2(
            Bucket=S3_BUCKET_NAME, Prefix=CALLNAME_CSV_PREFIX)
        for obj in resp.get("Contents", []):
            key = obj["Key"]
            if not key.lower().endswith(".csv"):
                continue
            size = obj.get("Size", 0)
            total_size += size
            files.append({
                "name": key.split("/")[-1],
                "size": size,
                "last_modified": obj["LastModified"].isoformat() if obj.get("LastModified") else None,
            })
    except Exception:
        pass
    return {
        "loaded": len(files) > 0,
        "rows": _callname_db_row_count,
        "file_count": len(files),
        "total_size": total_size,
        "files": files,
    }


@router.get("/callname/db-preview")
async def callname_db_preview(request: Request, limit: int = Query(50, ge=1, le=200)):
    """호출명칭 DB 미리보기 — S3 CSV에서 첫 N행 반환."""
    await _verify_auth(request)
    import csv as _csv_mod
    s3 = get_s3_client()
    result_files = []
    try:
        resp = s3.list_objects_v2(Bucket=S3_BUCKET_NAME, Prefix=CALLNAME_CSV_PREFIX)
        for obj in resp.get("Contents", []):
            key = obj["Key"]
            if not key.lower().endswith(".csv"):
                continue
            s3_obj = s3.get_object(Bucket=S3_BUCKET_NAME, Key=key)
            body_bytes = b""
            for chunk in s3_obj["Body"].iter_chunks(1024 * 1024):
                body_bytes = chunk
                break
            text = body_bytes.decode("utf-8", errors="replace")
            lines = text.split("\n")
            reader = _csv_mod.reader(lines)
            headers = []
            rows = []
            for i, row in enumerate(reader):
                if i == 0:
                    headers = row
                    continue
                if not any(row):
                    continue
                rows.append(row)
                if len(rows) >= limit:
                    break
            result_files.append({
                "name": key.split("/")[-1],
                "headers": headers,
                "rows": rows,
                "preview_count": len(rows),
            })
    except Exception as e:
        logger.error(f"호출명칭 DB 미리보기 실패: {e}")
        raise HTTPException(status_code=500, detail="호출명칭 DB 미리보기 실패")
    return {"files": result_files}


@router.post("/callname/upload-raw")
async def callname_upload_raw(request: Request, file: UploadFile = File(...)):
    """호출명칭 Excel → S3 멀티파트 스트리밍 (파싱 없음, 파일 전송만)."""
    await _verify_auth(request)
    _check_memory("호출명칭 업로드")
    _cleanup_callname_sessions()
    active = sum(1 for v in _callname_sessions.values() if v.get("status") in ("uploaded", "ready"))
    if active >= CALLNAME_MAX_SESSIONS:
        raise HTTPException(status_code=429, detail=f"동시 세션 초과 (최대 {CALLNAME_MAX_SESSIONS})")

    filename = file.filename or "unknown.xlsx"
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ("xlsx", "xls"):
        raise HTTPException(status_code=400, detail="xlsx 또는 xls 파일만 가능합니다.")

    safe_name = re.sub(r"[^\w\-_\.]", "_", filename)
    upload_id = str(uuid.uuid4())
    s3_key = f"callname-temp/{upload_id}/{safe_name}"
    content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if ext == "xlsx" else "application/vnd.ms-excel"
    )
    s3 = get_s3_client()
    mpu_upload_id: Optional[str] = None

    try:
        mpu = await asyncio.to_thread(
            lambda: s3.create_multipart_upload(
                Bucket=S3_BUCKET_NAME, Key=s3_key, ContentType=content_type
            )
        )
        mpu_upload_id = mpu["UploadId"]

        PART_SIZE = 8 * 1024 * 1024
        buf = b""
        parts: list = []
        part_number = 1
        total_size = 0

        while True:
            chunk = await file.read(PART_SIZE)
            if not chunk:
                break
            total_size += len(chunk)
            if total_size > MAX_DS_UPLOAD_SIZE:
                raise HTTPException(status_code=413, detail="파일 크기 초과 (200MB)")
            buf += chunk
            while len(buf) >= PART_SIZE:
                part_data, buf = buf[:PART_SIZE], buf[PART_SIZE:]
                pn = part_number
                resp = await asyncio.to_thread(
                    lambda pd=part_data, n=pn: s3.upload_part(
                        Bucket=S3_BUCKET_NAME, Key=s3_key,
                        UploadId=mpu_upload_id, PartNumber=n, Body=pd,
                    )
                )
                parts.append({"PartNumber": pn, "ETag": resp["ETag"]})
                part_number += 1

        if buf:
            pn = part_number
            resp = await asyncio.to_thread(
                lambda pd=buf, n=pn: s3.upload_part(
                    Bucket=S3_BUCKET_NAME, Key=s3_key,
                    UploadId=mpu_upload_id, PartNumber=n, Body=pd,
                )
            )
            parts.append({"PartNumber": pn, "ETag": resp["ETag"]})

        if not parts:
            raise ValueError("업로드된 데이터가 없습니다")

        await asyncio.to_thread(
            lambda: s3.complete_multipart_upload(
                Bucket=S3_BUCKET_NAME, Key=s3_key, UploadId=mpu_upload_id,
                MultipartUpload={"Parts": parts},
            )
        )
        logger.info(f"callname upload-raw: {s3_key} ({len(parts)} parts, {total_size // 1024}KB)")
        return {
            "success": True,
            "s3Key": s3_key,
            "uploadId": upload_id,
            "filename": filename,
            "ext": ext,
        }

    except HTTPException:
        raise
    except Exception as e:
        if mpu_upload_id:
            try:
                await asyncio.to_thread(
                    lambda: s3.abort_multipart_upload(
                        Bucket=S3_BUCKET_NAME, Key=s3_key, UploadId=mpu_upload_id,
                    )
                )
            except Exception:
                pass
        logger.error(f"호출명칭 upload-raw 실패: {e}")
        raise HTTPException(status_code=500, detail="서버 내부 오류")


@router.post("/callname/upload-complete")
async def callname_upload_complete(request: Request):
    """S3 업로드 완료 후 경량 파싱 — 컬럼 감지 + 행 수 집계."""
    await _verify_auth(request)
    _check_memory("호출명칭 파싱")

    body = await request.json()
    upload_id = body.get("uploadId")
    s3_key = body.get("s3Key")
    filename = body.get("filename", "unknown.xlsx")
    ext = body.get("ext", filename.rsplit(".", 1)[-1].lower())

    if not upload_id or not s3_key:
        raise HTTPException(status_code=400, detail="uploadId, s3Key 필수")

    if upload_id in _callname_sessions:
        sess = _callname_sessions[upload_id]
        return {
            "upload_id": upload_id,
            "filename": sess.get("filename", filename),
            **{k: sess.get(k) for k in ("total_rows", "columns", "callname_col",
                                          "tongsi_col", "zpwina_col", "zpwino_col",
                                          "filtered_rows")},
            "detected_callname_col": sess.get("callname_col"),
            "detected_tongsi_col": sess.get("tongsi_col"),
            "detected_zpwina_col": sess.get("zpwina_col"),
            "detected_zpwino_col": sess.get("zpwino_col"),
        }

    try:
        tmp_path = await asyncio.to_thread(_s3_to_tempfile, s3_key, f".{ext}")

        def _parse_lightweight_from_s3():
            if ext == "xlsx":
                fast = _parse_xlsx_header_fast(tmp_path)
                columns = fast["columns"]
                total_rows = fast["total_rows"]
                tongsi_col = _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS)
                filtered_rows = total_rows
            else:
                import xlrd as _xlrd
                wb = _xlrd.open_workbook(tmp_path)
                ws = wb.sheet_by_index(0)
                columns = [str(ws.cell_value(0, c)) for c in range(ws.ncols)]
                total_rows = ws.nrows - 1
                tongsi_col = _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS)
                filtered_rows = total_rows
                wb.release_resources()

            callname_col = _detect_column(columns, CALLNAME_POSSIBLE_CALLNAME_COLS)
            zpwina_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINA_COLS)
            zpwino_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINO_COLS)

            return {
                "total_rows": total_rows, "columns": columns,
                "callname_col": callname_col, "tongsi_col": tongsi_col,
                "zpwina_col": zpwina_col, "zpwino_col": zpwino_col,
                "filtered_rows": filtered_rows,
            }

        info = await asyncio.to_thread(_parse_lightweight_from_s3)

        _callname_sessions[upload_id] = {
            "filename": filename,
            "s3_temp_key": s3_key,
            "ext": ext,
            "status": "uploaded",
            "created_at_ts": _time_mod.time(),
            "cached_xlsx_path": tmp_path,
            "analysis_status": "pending",
            **info,
        }

        asyncio.get_event_loop().run_in_executor(None, _analyze_callname_bg, upload_id)

        return {
            "upload_id": upload_id,
            "filename": filename,
            **info,
            "detected_callname_col": info["callname_col"],
            "detected_tongsi_col": info["tongsi_col"],
            "detected_zpwina_col": info["zpwina_col"],
            "detected_zpwino_col": info["zpwino_col"],
        }
    except HTTPException:
        raise
    except Exception as e:
        if 'tmp_path' in dir():
            try:
                os.remove(tmp_path)
            except Exception:
                pass
        logger.error(f"호출명칭 upload-complete 실패: {e}")
        raise HTTPException(status_code=500, detail="서버 내부 오류")


@router.get("/callname/upload/{upload_id}/analysis")
async def callname_analysis_status(upload_id: str, request: Request):
    """백그라운드 분석 상태 조회."""
    await _verify_auth(request)
    sess = _callname_sessions.get(upload_id)
    if not sess:
        raise HTTPException(status_code=404, detail="세션 없음")
    status = sess.get("analysis_status", "pending")
    result = {"status": status}
    if status == "complete":
        result["filtered_rows"] = sess.get("filtered_rows", 0)
        result["target_callnames"] = sess.get("target_callnames", 0)
        result["total_rows"] = sess.get("total_rows", 0)
        result["columns"] = sess.get("columns", [])
        result["detected_callname_col"] = sess.get("callname_col")
        result["detected_tongsi_col"] = sess.get("tongsi_col")
        result["detected_zpwina_col"] = sess.get("zpwina_col")
        result["detected_zpwino_col"] = sess.get("zpwino_col")
    return result


@router.post("/callname/upload")
async def callname_upload(request: Request, file: UploadFile = File(...)):
    """Excel 업로드 → S3 임시저장 + 경량 컬럼 감지 (소용량 fallback)."""
    await _verify_auth(request)
    _check_memory("호출명칭 Excel 업로드")

    _cleanup_callname_sessions()
    active = sum(1 for v in _callname_sessions.values() if v.get("status") in ("uploaded", "ready"))
    if active >= CALLNAME_MAX_SESSIONS:
        raise HTTPException(status_code=429, detail=f"동시 세션 초과 (최대 {CALLNAME_MAX_SESSIONS})")

    filename = file.filename or "unknown.xlsx"
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ("xlsx", "xls"):
        raise HTTPException(status_code=400, detail="xlsx 또는 xls 파일만 가능합니다.")

    tmp = _tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
    tmp_path = tmp.name
    file_size = 0
    try:
        while True:
            chunk = await file.read(4 * 1024 * 1024)
            if not chunk:
                break
            file_size += len(chunk)
            if file_size > MAX_DS_UPLOAD_SIZE:
                tmp.close()
                os.remove(tmp_path)
                raise HTTPException(status_code=413, detail="파일 크기 초과 (200MB)")
            tmp.write(chunk)
        tmp.close()

        safe_name = re.sub(r"[^\w\-_\.]", "_", filename)
        upload_id = str(uuid.uuid4())
        s3_key = f"callname-temp/{upload_id}/{safe_name}"

        with open(tmp_path, "rb") as f:
            await asyncio.to_thread(
                lambda: get_s3_client().put_object(
                    Bucket=S3_BUCKET_NAME, Key=s3_key, Body=f,
                    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    if ext == "xlsx" else "application/vnd.ms-excel",
                )
            )

        def _parse_lightweight():
            if ext == "xlsx":
                fast = _parse_xlsx_header_fast(tmp_path)
                columns = fast["columns"]
                total_rows = fast["total_rows"]
                tongsi_col = _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS)
                filtered_rows = total_rows
            else:
                import xlrd as _xlrd
                wb = _xlrd.open_workbook(tmp_path)
                ws = wb.sheet_by_index(0)
                columns = [str(ws.cell_value(0, c)) for c in range(ws.ncols)]
                total_rows = ws.nrows - 1
                tongsi_col = _detect_column(columns, CALLNAME_POSSIBLE_TONGSI_COLS)
                filtered_rows = total_rows
                wb.release_resources()

            callname_col = _detect_column(columns, CALLNAME_POSSIBLE_CALLNAME_COLS)
            zpwina_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINA_COLS)
            zpwino_col = _detect_column(columns, CALLNAME_POSSIBLE_ZPWINO_COLS)

            return {
                "total_rows": total_rows, "columns": columns,
                "callname_col": callname_col, "tongsi_col": tongsi_col,
                "zpwina_col": zpwina_col, "zpwino_col": zpwino_col,
                "filtered_rows": filtered_rows,
            }

        info = await asyncio.to_thread(_parse_lightweight)

        _callname_sessions[upload_id] = {
            "filename": filename,
            "s3_temp_key": s3_key,
            "ext": ext,
            "status": "uploaded",
            "created_at_ts": _time_mod.time(),
            "cached_xlsx_path": tmp_path,
            "analysis_status": "pending",
            **info,
        }

        asyncio.get_event_loop().run_in_executor(None, _analyze_callname_bg, upload_id)

        return {
            "upload_id": upload_id,
            "filename": filename,
            **info,
            "detected_callname_col": info["callname_col"],
            "detected_tongsi_col": info["tongsi_col"],
            "detected_zpwina_col": info["zpwina_col"],
            "detected_zpwino_col": info["zpwino_col"],
        }
    except HTTPException:
        raise
    except Exception as e:
        try:
            os.remove(tmp_path)
        except Exception:
            pass
        logger.error(f"호출명칭 업로드 실패: {e}")
        raise HTTPException(status_code=500, detail="서버 내부 오류")


@router.post("/callname/upload/{upload_id}/column-values")
async def callname_column_values(upload_id: str, request: Request):
    """컬럼 고유값 조회."""
    await _verify_auth(request)
    if upload_id not in _callname_sessions:
        raise HTTPException(status_code=404, detail="세션 없음")
    sess = _callname_sessions[upload_id]
    if sess.get("status") != "uploaded":
        raise HTTPException(status_code=400, detail="이미 처리 시작됨")

    body = await request.json()
    col = body.get("column")
    if not col or col not in sess.get("columns", []):
        raise HTTPException(status_code=400, detail=f"'{col}' 컬럼 없음")

    column_stats = sess.get("column_stats", {})
    if col in column_stats:
        return {"column": col, "values": column_stats[col]}

    def _calc():
        from collections import Counter
        columns = sess.get("columns", [])
        col_idx = columns.index(col) if col in columns else -1
        if col_idx < 0:
            return []

        cached_xlsx = sess.get("cached_xlsx_path")
        cached_ss = sess.get("cached_ss_path")
        cached_ss_off = sess.get("cached_ss_offsets")

        if cached_xlsx and os.path.exists(cached_xlsx):
            tmp_path = cached_xlsx
            need_cleanup = False
        else:
            tmp_path = _s3_to_tempfile(sess["s3_temp_key"], suffix=f".{sess['ext']}")
            need_cleanup = True
            cached_ss = None
            cached_ss_off = None
        try:
            counter = Counter()
            if sess["ext"] == "xlsx":
                for rn, vals in _iter_xlsx_rows_light(
                        tmp_path, ss_cache_path=cached_ss, ss_offsets_bytes=cached_ss_off):
                    if rn == 0:
                        continue
                    v = vals[col_idx] if col_idx < len(vals) else ""
                    if v:
                        counter[v] += 1
                _release_memory()
            else:
                import xlrd as _xlrd
                xls_book = _xlrd.open_workbook(tmp_path)
                ws = xls_book.sheet_by_index(0)
                for r in range(1, ws.nrows):
                    v = str(ws.cell_value(r, col_idx)) if col_idx < ws.ncols else ""
                    if v:
                        counter[v] += 1
                xls_book.release_resources()
            return [{"value": v, "count": c} for v, c in counter.most_common(100)]
        finally:
            if need_cleanup:
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass

    values = await asyncio.to_thread(_calc)
    return {"column": col, "values": values}


@router.post("/callname/upload/{upload_id}/preview")
async def callname_preview(upload_id: str, request: Request):
    """필터 미리보기 — 분석 시 캐시된 tongsi 빈 행 데이터로 즉시 계산."""
    await _verify_auth(request)
    if upload_id not in _callname_sessions:
        raise HTTPException(status_code=404, detail="세션 없음")
    sess = _callname_sessions[upload_id]
    if sess.get("status") != "uploaded":
        raise HTTPException(status_code=400, detail="이미 처리 시작됨")

    body = await request.json()
    filters = body.get("filters", {})
    callname_col = sess.get("callname_col")

    if not filters and sess.get("analysis_status") == "complete":
        return {
            "filtered_rows": sess.get("filtered_rows", 0),
            "target_callnames": sess.get("target_callnames", 0),
        }

    cached_rows = sess.get("filter_cache_rows")
    if cached_rows is not None:
        columns = sess.get("columns", [])
        callname_idx = columns.index(callname_col) if callname_col and callname_col in columns else -1

        filter_col_indices = {}
        if filters:
            for c, vals in filters.items():
                if c in columns and vals:
                    filter_col_indices[columns.index(c)] = set(str(v) for v in vals)

        filtered_rows = 0
        callname_set = set()
        for row_vals in cached_rows:
            passed = True
            for ci, allowed in filter_col_indices.items():
                if ci < len(row_vals) and row_vals[ci] not in allowed:
                    passed = False
                    break
            if not passed:
                continue
            filtered_rows += 1
            if 0 <= callname_idx < len(row_vals) and row_vals[callname_idx].strip():
                callname_set.add(row_vals[callname_idx].strip())

        return {"filtered_rows": filtered_rows, "target_callnames": len(callname_set)}

    return {
        "filtered_rows": sess.get("filtered_rows", 0),
        "target_callnames": sess.get("target_callnames", 0),
    }


@router.post("/callname/process")
async def callname_process(request: Request):
    """매칭 시작 — filter_cache_rows 캐시 활용."""
    await _verify_auth(request)
    _check_rate_limit(request, "callname_process", 3, 60)

    body = await request.json()
    upload_id = body.get("upload_id")
    filters = body.get("filters", {})

    if not upload_id or upload_id not in _callname_sessions:
        raise HTTPException(status_code=404, detail="세션 없음")
    sess = _callname_sessions[upload_id]
    if sess.get("status") != "uploaded":
        raise HTTPException(status_code=400, detail="이미 처리 시작됨")

    zpwina_col = sess.get("zpwina_col")
    zpwino_col = sess.get("zpwino_col")

    if not zpwina_col and not zpwino_col:
        raise HTTPException(status_code=400, detail="zpwina/zpwino 컬럼 없음")

    columns = sess.get("columns", [])
    cached_rows = sess.get("filter_cache_rows")
    if cached_rows is None:
        raise HTTPException(status_code=400, detail="분석 미완료 — 잠시 후 다시 시도해주세요.")

    zpwina_idx = columns.index(zpwina_col) if zpwina_col and zpwina_col in columns else -1
    zpwino_idx = columns.index(zpwino_col) if zpwino_col and zpwino_col in columns else -1

    filter_col_indices = {}
    if filters:
        for c, vals in filters.items():
            if c in columns and vals:
                filter_col_indices[columns.index(c)] = set(str(v) for v in vals)

    cached_row_indices = sess.get("filter_cache_row_indices", [])
    zpwina_set = set()
    zpwino_set = set()
    original_row_indices = []
    row_zpwina_list = []
    row_zpwino_list = []

    for i, row_vals in enumerate(cached_rows):
        passed = True
        for ci, allowed in filter_col_indices.items():
            if ci < len(row_vals) and row_vals[ci] not in allowed:
                passed = False
                break
        if not passed:
            continue

        excel_row = cached_row_indices[i] if i < len(cached_row_indices) else (i + 2)
        original_row_indices.append(excel_row)
        za = row_vals[zpwina_idx] if 0 <= zpwina_idx < len(row_vals) else ""
        zo = row_vals[zpwino_idx] if 0 <= zpwino_idx < len(row_vals) else ""
        row_zpwina_list.append(za)
        row_zpwino_list.append(zo)
        if za:
            zpwina_set.add(za)
        if zo:
            zpwino_set.add(zo)

    _log_mem("callname_process 완료 (캐시 활용)")

    total_values = len(zpwina_set | zpwino_set)
    if total_values == 0:
        raise HTTPException(status_code=400, detail="매칭 대상 값이 없습니다.")

    process_id = str(uuid.uuid4())
    _callname_sessions[process_id] = {
        "s3_temp_key": sess["s3_temp_key"],
        "ext": sess["ext"],
        "original_row_indices": original_row_indices,
        "row_zpwina_list": row_zpwina_list,
        "row_zpwino_list": row_zpwino_list,
        "zpwina_col": zpwina_col,
        "zpwino_col": zpwino_col,
        "zpwina_values": list(zpwina_set),
        "zpwino_values": list(zpwino_set),
        "filename": sess["filename"],
        "columns": columns,
        "cached_xlsx_path": sess.get("cached_xlsx_path"),
        "status": "ready",
        "created_at_ts": _time_mod.time(),
    }
    del _callname_sessions[upload_id]
    _release_memory()

    return {
        "process_id": process_id,
        "total_values": total_values,
        "total_rows": len(row_zpwina_list),
    }


@router.get("/callname/process/{process_id}/stream")
async def callname_stream(process_id: str, request: Request):
    """SSE 스트리밍 — 호출명칭 매칭 결과 스트림."""
    await _verify_auth(request)
    _check_memory("호출명칭 매칭")
    if process_id not in _callname_sessions:
        raise HTTPException(status_code=404, detail="세션 없음")

    sess = _callname_sessions[process_id]

    def _generate():
        import json
        try:
            zpwina_values = sess["zpwina_values"]
            zpwino_values = sess["zpwino_values"]
            filename = sess["filename"]
            s3_temp_key = sess["s3_temp_key"]
            ext = sess["ext"]
            original_row_indices = sess.get("original_row_indices", [])

            total_values = len(set(zpwina_values + zpwino_values))
            total_rows = len(original_row_indices)

            _log_mem("stream 시작")
            yield f"data: {json.dumps({'type': 'progress', 'progress': 5, 'message': '파일 분석 완료', 'detail': f'{total_rows:,}행, {total_values:,}개 고유값'})}\n\n"

            _log_mem("1단계: DB 조회 시작")
            yield f"data: {json.dumps({'type': 'progress', 'progress': 10, 'message': 'DB 로드 + 6방향 교차 조회 중...'})}\n\n"

            db_data = _query_callname_db(zpwina_values, zpwino_values)
            db_count = len(db_data)
            _log_mem("1단계: DB 조회 완료")

            yield f"data: {json.dumps({'type': 'progress', 'progress': 40, 'message': 'DB 조회 완료', 'detail': f'{db_count:,}건 매칭됨'})}\n\n"

            _log_mem("2단계: 매칭 시작")
            yield f"data: {json.dumps({'type': 'progress', 'progress': 45, 'message': '매칭 데이터 준비 중...'})}\n\n"

            columns = sess.get("columns", [])

            db_fields = ["area_hdofc_nm", "ons_team_nm", "zpcode"]
            excel_col_map = {}
            for db_field, candidates in CALLNAME_DB_TO_EXCEL_MAP.items():
                detected = _detect_column(columns, candidates)
                excel_col_map[db_field] = detected if detected else candidates[0]
            target_excel_cols = [excel_col_map[f] for f in db_fields]

            row_zpwina_list = sess.get("row_zpwina_list", [])
            row_zpwino_list = sess.get("row_zpwino_list", [])
            cached_xlsx = sess.get("cached_xlsx_path")
            if cached_xlsx and os.path.exists(cached_xlsx):
                tmp_excel_path = cached_xlsx
                _log_mem("2단계: 캐시된 xlsx 재사용")
            else:
                _log_mem("2단계: S3 temp 다운로드 시작")
                tmp_excel_path = _s3_to_tempfile(s3_temp_key, suffix=f".{ext}")
                _log_mem("2단계: S3 temp 다운로드 완료")

            matched_count = 0
            zpwina_matched = 0
            zpwino_matched = 0
            cross_matched = 0
            row_data_map = {}

            for i, row_idx in enumerate(original_row_indices):
                excel_row_num = row_idx
                za = row_zpwina_list[i] if i < len(row_zpwina_list) else ""
                zo = row_zpwino_list[i] if i < len(row_zpwino_list) else ""
                hit = None
                if za:
                    hit = db_data.get(za)
                    if hit and hit.get("zpcode"):
                        row_data_map[excel_row_num] = [
                            hit.get("area_hdofc_nm", ""),
                            hit.get("ons_team_nm", ""),
                            hit.get("zpcode", ""),
                        ]
                        zpwina_matched += 1
                        continue
                if zo:
                    hit = db_data.get(zo)
                    if hit and hit.get("zpcode"):
                        row_data_map[excel_row_num] = [
                            hit.get("area_hdofc_nm", ""),
                            hit.get("ons_team_nm", ""),
                            hit.get("zpcode", ""),
                        ]
                        zpwino_matched += 1

            matched_count = len(row_data_map)
            del db_data, row_zpwina_list, row_zpwino_list
            _release_memory()
            _log_mem("2단계: 매칭 완료")

            yield f"data: {json.dumps({'type': 'progress', 'progress': 60, 'message': '매칭 완료', 'detail': f'{matched_count:,}/{total_rows:,}행 (zpwina:{zpwina_matched:,}, zpwino:{zpwino_matched:,})'})}\n\n"

            _log_mem("3단계: ZIP XML 시작")
            yield f"data: {json.dumps({'type': 'progress', 'progress': 65, 'message': 'Excel 파일 생성 중...'})}\n\n"

            tmp_output = _tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
            tmp_output_path = tmp_output.name
            tmp_output.close()

            try:
                with zipfile.ZipFile(tmp_excel_path, "r") as zin:
                    sheet_files = [f for f in zin.namelist() if "worksheets/sheet" in f]
                    sheet_path = sheet_files[0] if sheet_files else "xl/worksheets/sheet1.xml"

                    target_col_letters = []
                    for ecn in target_excel_cols:
                        if ecn in columns:
                            idx = columns.index(ecn) + 1
                            target_col_letters.append(get_column_letter(idx))
                        else:
                            target_col_letters.append(get_column_letter(len(columns) + 1 + len(target_col_letters)))

                    target_letters_set = set(target_col_letters)
                    cell_pattern = re.compile(r'(<c r="([A-Z]+)\d+"[^>]*(?:>.*?</c>|/>))', re.DOTALL)

                    tmp_sheet = _tempfile.NamedTemporaryFile(delete=False, suffix=".xml", mode="w", encoding="utf-8")
                    tmp_sheet_path = tmp_sheet.name

                    with zin.open(sheet_path) as sheet_stream:
                        buffer = ""
                        CHUNK_SIZE = 512 * 1024
                        while True:
                            raw = sheet_stream.read(CHUNK_SIZE)
                            if not raw:
                                break
                            buffer += raw.decode("utf-8", errors="replace")

                            while "</row>" in buffer:
                                row_end = buffer.index("</row>") + 6
                                row_section = buffer[:row_end]
                                buffer = buffer[row_end:]

                                r_pos = row_section.find('<row r="')
                                if r_pos == -1:
                                    tmp_sheet.write(row_section)
                                    continue

                                r_start = r_pos + 8
                                r_end_q = row_section.index('"', r_start)
                                row_num = int(row_section[r_start:r_end_q])
                                vals = row_data_map.get(row_num)

                                if vals is None:
                                    tmp_sheet.write(row_section)
                                else:
                                    part = row_section[:-6]
                                    row_tag_start = part.find('<row r="')
                                    row_tag_end = part.index(">", row_tag_start) + 1
                                    before_row = part[:row_tag_start]
                                    row_tag = part[row_tag_start:row_tag_end]
                                    after_row_tag = part[row_tag_end:]

                                    cell_dict = {}
                                    for cell_match in cell_pattern.finditer(after_row_tag):
                                        full_cell = cell_match.group(1)
                                        cl = cell_match.group(2)
                                        if cl not in target_letters_set:
                                            cell_dict[cl] = full_cell

                                    for i, val in enumerate(vals):
                                        cl = target_col_letters[i]
                                        safe_val = str(val).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
                                        if safe_val:
                                            cell_dict[cl] = f'<c r="{cl}{row_num}" t="inlineStr"><is><t>{safe_val}</t></is></c>'

                                    sorted_cells = sorted(cell_dict.items(),
                                        key=lambda x: column_index_from_string(x[0]))
                                    tmp_sheet.write(before_row)
                                    tmp_sheet.write(row_tag)
                                    for _, xml in sorted_cells:
                                        tmp_sheet.write(xml)
                                    tmp_sheet.write("</row>")

                        if buffer:
                            tmp_sheet.write(buffer)

                    tmp_sheet.close()
                    del row_data_map
                    _release_memory()

                    _log_mem("3단계: XML 스트리밍 완료")
                    yield f"data: {json.dumps({'type': 'progress', 'progress': 85, 'message': 'ZIP 재조립 중...'})}\n\n"

                    with zipfile.ZipFile(tmp_output_path, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as zout:
                        for item in zin.infolist():
                            if item.filename == sheet_path:
                                zout.write(tmp_sheet_path, item.filename)
                            else:
                                with zin.open(item.filename) as src, zout.open(item, "w") as dst:
                                    while True:
                                        chunk = src.read(512 * 1024)
                                        if not chunk:
                                            break
                                        dst.write(chunk)

                try:
                    os.remove(tmp_sheet_path)
                except Exception:
                    pass
                _release_memory()

                _log_mem("4단계: ZIP 재조립 완료")
                yield f"data: {json.dumps({'type': 'progress', 'progress': 92, 'message': 'S3 업로드 중...'})}\n\n"

                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                base_name = filename.rsplit(".", 1)[0]
                output_filename = f"{base_name}_matched_{timestamp}.xlsx"
                s3_result_key = f"callname-results/{process_id}/{output_filename}"

                with open(tmp_output_path, "rb") as f:
                    get_s3_client().upload_fileobj(
                        f, S3_BUCKET_NAME, s3_result_key,
                        ExtraArgs={"ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
                    )

                sess["s3_result_key"] = s3_result_key
                sess["output_filename"] = output_filename
                sess["status"] = "completed"

                for _drop_key in ("original_row_indices", "row_zpwina_list", "row_zpwino_list",
                                  "zpwina_values", "zpwino_values", "columns",
                                  "cached_xlsx_path", "column_stats", "cached_ss_offsets"):
                    sess.pop(_drop_key, None)

                yield f"data: {json.dumps({'type': 'complete', 'progress': 100, 'message': f'완료! (매칭: {matched_count:,}/{total_rows:,}건)', 'matched': matched_count, 'total': total_rows, 'zpwina_matched': zpwina_matched, 'zpwino_matched': zpwino_matched, 'cross_matched': cross_matched})}\n\n"

            finally:
                for _p in [tmp_output_path, tmp_excel_path]:
                    try:
                        os.remove(_p)
                    except Exception:
                        pass
                _release_memory()
                _log_mem("stream 종료 (정리 완료)")

        except Exception as e:
            logger.exception(f"호출명칭 매칭 스트림 오류: {e}")
            try:
                os.remove(tmp_excel_path)
            except Exception:
                pass
            _release_memory()
            import json
            yield f"data: {json.dumps({'type': 'error', 'message': '서버 내부 오류'})}\n\n"

    return StreamingResponse(
        _generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.get("/callname/process/{process_id}/download")
async def callname_download(process_id: str, request: Request):
    """매칭 결과 Excel 다운로드 (S3 presign URL)."""
    await _verify_auth(request)
    if process_id not in _callname_sessions:
        raise HTTPException(status_code=404, detail="세션 없음")

    data = _callname_sessions[process_id]
    if data.get("status") != "completed":
        raise HTTPException(status_code=400, detail="처리 미완료")

    s3_key = data.get("s3_result_key")
    output_filename = data.get("output_filename", "result.xlsx")

    if not s3_key:
        raise HTTPException(status_code=500, detail="결과 파일 없음")

    try:
        encoded_filename = quote(output_filename, safe="")
        url = get_s3_client().generate_presigned_url(
            "get_object",
            Params={
                "Bucket": S3_BUCKET_NAME,
                "Key": s3_key,
                "ResponseContentDisposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            },
            ExpiresIn=600,
        )
        return {"url": url, "filename": output_filename}
    except Exception as e:
        logger.error(f"호출명칭 결과 다운로드 실패: {e}")
        raise HTTPException(status_code=500, detail="서버 내부 오류")


@router.get("/callname/sample-template")
async def callname_sample_template_list(request: Request):
    """호출명칭 sample 양식 목록 조회."""
    await _verify_auth(request)
    files = []
    try:
        resp = get_s3_client().list_objects_v2(
            Bucket=S3_BUCKET_NAME, Prefix=CALLNAME_SAMPLE_PREFIX)
        for obj in resp.get("Contents", []):
            key = obj["Key"]
            name = key[len(CALLNAME_SAMPLE_PREFIX):]
            if not name:
                continue
            if not name.lower().endswith(CALLNAME_SAMPLE_ALLOWED_EXTS):
                continue
            files.append({
                "name": name,
                "size": obj.get("Size", 0),
                "last_modified": obj["LastModified"].isoformat()
                    if obj.get("LastModified") else None,
            })
    except Exception as e:
        logger.error(f"호출명칭 sample 목록 조회 실패: {e}")
        raise HTTPException(status_code=500, detail="목록 조회 실패")
    files.sort(key=lambda x: x.get("last_modified") or "", reverse=True)
    return {"files": files}


@router.post("/callname/sample-template")
async def callname_sample_template_upload(
    request: Request,
    file: UploadFile = File(...),
):
    """호출명칭 sample 양식 업로드 (admin 전용)."""
    await _require_role(request, {"admin"})

    raw_name = file.filename or ""
    base_name = os.path.basename(raw_name)
    if not base_name:
        raise HTTPException(status_code=400, detail="파일명이 비어있습니다.")
    ext = base_name.rsplit(".", 1)[-1].lower() if "." in base_name else ""
    if ext not in CALLNAME_SAMPLE_ALLOWED_EXTS:
        raise HTTPException(status_code=400, detail="xlsx 또는 xls 파일만 가능합니다.")

    safe_name = re.sub(r"[^\w\-\.가-힣]", "_", base_name)
    if not safe_name or safe_name.startswith("."):
        raise HTTPException(status_code=400, detail="유효한 파일명이 아닙니다.")

    total = 0
    chunks: list = []
    while True:
        chunk = await file.read(2 * 1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > CALLNAME_SAMPLE_MAX_SIZE:
            raise HTTPException(status_code=413, detail="파일 크기 초과 (최대 20MB)")
        chunks.append(chunk)
    body_bytes = b"".join(chunks)

    s3_key = f"{CALLNAME_SAMPLE_PREFIX}{safe_name}"
    content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if ext == "xlsx" else "application/vnd.ms-excel"
    )
    try:
        await asyncio.to_thread(
            lambda: get_s3_client().put_object(
                Bucket=S3_BUCKET_NAME, Key=s3_key,
                Body=body_bytes, ContentType=content_type,
            )
        )
    except Exception as e:
        logger.error(f"호출명칭 sample 업로드 실패: {e}")
        raise HTTPException(status_code=500, detail="업로드 실패")

    try:
        empno = await _verify_auth(request)
        await asyncio.to_thread(
            _record_audit_log_sync,
            "upload", "callname_sample", safe_name, empno,
            {"size": total},
        )
    except Exception:
        pass

    return {"success": True, "name": safe_name, "size": total}


@router.get("/callname/sample-template/download")
async def callname_sample_template_download(
    request: Request,
    name: str = Query(..., description="파일명"),
):
    """호출명칭 sample 양식 다운로드 (presign URL)."""
    await _verify_auth(request)
    safe_name = os.path.basename(name or "")
    if not safe_name or safe_name.startswith("."):
        raise HTTPException(status_code=400, detail="유효한 파일명이 아닙니다.")
    if not safe_name.lower().endswith(CALLNAME_SAMPLE_ALLOWED_EXTS):
        raise HTTPException(status_code=400, detail="허용되지 않은 파일 형식")

    s3_key = f"{CALLNAME_SAMPLE_PREFIX}{safe_name}"
    s3 = get_s3_client()
    try:
        await asyncio.to_thread(
            lambda: s3.head_object(Bucket=S3_BUCKET_NAME, Key=s3_key))
    except Exception:
        raise HTTPException(status_code=404, detail="파일 없음")

    try:
        encoded = quote(safe_name, safe="")
        url = s3.generate_presigned_url(
            "get_object",
            Params={
                "Bucket": S3_BUCKET_NAME,
                "Key": s3_key,
                "ResponseContentDisposition":
                    f"attachment; filename*=UTF-8''{encoded}",
            },
            ExpiresIn=600,
        )
        return {"url": url, "filename": safe_name}
    except Exception as e:
        logger.error(f"호출명칭 sample 다운로드 URL 생성 실패: {e}")
        raise HTTPException(status_code=500, detail="다운로드 실패")


@router.delete("/callname/sample-template")
async def callname_sample_template_delete(
    request: Request,
    name: str = Query(..., description="파일명"),
):
    """호출명칭 sample 양식 삭제 (admin 전용)."""
    empno = await _require_role(request, {"admin"})
    safe_name = os.path.basename(name or "")
    if not safe_name or safe_name.startswith("."):
        raise HTTPException(status_code=400, detail="유효한 파일명이 아닙니다.")

    s3_key = f"{CALLNAME_SAMPLE_PREFIX}{safe_name}"
    try:
        await asyncio.to_thread(
            lambda: get_s3_client().delete_object(
                Bucket=S3_BUCKET_NAME, Key=s3_key))
    except Exception as e:
        logger.error(f"호출명칭 sample 삭제 실패: {e}")
        raise HTTPException(status_code=500, detail="삭제 실패")

    try:
        await asyncio.to_thread(
            _record_audit_log_sync,
            "delete", "callname_sample", safe_name, empno, None,
        )
    except Exception:
        pass

    return {"success": True}
