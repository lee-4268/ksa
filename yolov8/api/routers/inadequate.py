"""
inadequate - 부적합 관리 엔드포인트

담당 도메인: 부적합 국소 관리/동기화/엑셀 내보내기
주요 의존성: core.auth, core.config
엔드포인트:
    POST /inadequate/sync
    GET  /inadequate/list
    PUT  /inadequate/update
    GET  /inadequate/stats
    GET  /inadequate/export-xlsx
"""

import asyncio
import io
import logging
import re
import sqlite3
from datetime import datetime, timedelta, timezone
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import Response

from core.auth import _verify_auth, _get_user_role_sync
from core.config import _INSP_DB
from schemas.models import InadequateUpdateReq

router = APIRouter(tags=["inadequate"])
logger = logging.getLogger(__name__)


# ── 부적합 유형 분류 ──────────────────────────────────────────
# 근거는 불합격상세 컬럼이다. 여기에는 결과장 Y열(부적합내용)이 들어간다
# (inadequate_sync 주석 참조). 불합격내용·원문 불합격상세는 대부분 비어 있어
# 쓸 수 없다 — 2026 실측 58건 중 불합격내용 45건, X열 원문 48건이 공란.
#
# 전량(1,702건) 실측 결과: 공용화 1,355 / 주소 194 / 미분류 153.
#   미분류 153 중 149 가 공란이라 실제 누락은 4건이다.
#     '서류 : 프레임, 현재 : 간이폴 및 비기준 설치대' x3 — 설치대 건이라 두 유형 중
#        어디에도 속하지 않는다. 주소로 넣으면 안 되므로 미분류로 둔다.
#     '황지동 산69 --> 통동 산69' x1 — 산번지. 아래 '산\d' 패턴으로 잡는다.
#
# kca-be 에도 같은 규칙이 있다(apps/routers/inadequate.py). 한쪽만 고치면
# 두 화면의 집계가 어긋나므로 반드시 같이 바꿀 것.
_INAD_SHARE_KEYS = ("공용", "개별", "환경")
_INAD_ADDR_KEYS = ("설치장소", "주소")
# 지번 표기 세 갈래:
#   '안현동 204'  — 행정구역 접미사 + 숫자
#   '산 90-48'    — 산번지('황지동 산69' 처럼 접미사와 숫자 사이에 끼기도 한다)
#   '803-16'      — 본번-부번
_RE_INAD_JIBUN = re.compile(
    r"[가-힣]+(?:동|리|로|길|읍|면|가)\s*\d|산\s*\d|\d+\s*-\s*\d")


def _inad_kind(detail: str) -> str:
    """부적합내용 → '공용화' | '주소' | '미분류'.

    공용화 키워드를 먼저 본다. '환경 -> 개별'처럼 두 축이 겹쳐 보이는 값이
    있는데 이는 공용화 구분 변경이지 주소 건이 아니다.
    """
    s = (detail or "").strip()
    if not s:
        return "미분류"
    if any(k in s for k in _INAD_SHARE_KEYS):
        return "공용화"
    if any(k in s for k in _INAD_ADDR_KEYS) or _RE_INAD_JIBUN.search(s):
        return "주소"
    return "미분류"


@router.post("/inadequate/sync")
async def inadequate_sync(request: Request, year: int = Query(...)):
    """실적 데이터에서 부적합 국소 동기화."""
    empno = await _verify_auth(request)
    role = await asyncio.to_thread(_get_user_role_sync, empno)
    if role not in {"admin", "manager"}:
        raise HTTPException(403, "관리자/매니저만 가능")

    def _sync():
        conn = sqlite3.connect(_INSP_DB, timeout=60)
        conn.row_factory = sqlite3.Row
        # V12 결과장부터 부적합 판정은 V열('부적합') 전용 열 기준.
        # (기존 성능서류 기준은 불합격 유형 열이라 대상이 어긋나 있었다.
        #  V열 마킹은 결과장 작성자가 다음 업로드부터 채우기로 함 — 초기엔 건수가 적을 수 있다.)
        rows = conn.execute(
            "SELECT * FROM inspection_results_raw WHERE year=? AND 부적합='부적합'",
            (year,)
        ).fetchall()
        count = 0
        for r in rows:
            검사일자_raw = r['검사일자'] or ''
            검사일자 = 검사일자_raw
            시정기한 = ''
            if 검사일자_raw:
                try:
                    dt = None
                    s = str(검사일자_raw).strip()
                    try:
                        serial = float(s)
                        if 40000 < serial < 60000:
                            dt = datetime(1899, 12, 30) + timedelta(days=int(serial))
                            검사일자 = dt.strftime('%Y-%m-%d')
                    except (ValueError, TypeError):
                        pass
                    if dt is None:
                        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%Y/%m/%d'):
                            try:
                                dt = datetime.strptime(s.split('.')[0].strip(), fmt)
                                검사일자 = dt.strftime('%Y-%m-%d')
                                break
                            except Exception:
                                pass
                    if dt:
                        month = dt.month + 6
                        year_add = (month - 1) // 12
                        month = ((month - 1) % 12) + 1
                        시정기한 = dt.replace(year=dt.year + year_add, month=month).strftime('%Y-%m-%d')
                except Exception:
                    pass
            # 불합격상세 자리에 Y열(부적합내용 — 짧은 사유 분류값)을 담는다.
            # kca 미러(/inadequate/sync-export)도 이 컬럼을 그대로 내보내므로
            # kca 는 코드 변경 없이 수신한다. X(불합격상세사유)는 여러 줄 서술이라
            # 여기 폴백으로 섞지 않는다 — 목록 표시·유형별 집계가 깨진다.
            # X 원문은 raw.불합격상세에 그대로 남아 결과 화면 특이사항에서 조회 가능.
            try:
                _사유 = (r['부적합내용'] or '').strip()
            except (IndexError, KeyError):
                _사유 = ''
            conn.execute('''INSERT OR IGNORE INTO inadequate_management
                (year, 허가번호, 통합시설코드, 호출명칭, 주소, skt본부, region, ons팀,
                 검사일자, 시정기한, 불합격내용, 불합격상세)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?)''',
                (year, r['허가번호'], r['통합시설코드'], r['호출명칭'], r['주소'],
                 r['skt본부'], r['region'], r['ons팀'],
                 검사일자, 시정기한, r['불합격내용'] or '', _사유))
            count += 1
        conn.commit()
        total = conn.execute("SELECT COUNT(*) FROM inadequate_management WHERE year=?", (year,)).fetchone()[0]
        conn.close()
        return {"synced": count, "total": total}

    return await asyncio.to_thread(_sync)


@router.post("/inadequate/sync-export")
async def inadequate_sync_export(request: Request):
    """부적합 전량 JSON — kca 미러링 원본. 읽기 전용 (ksa 데이터 변경 없음).

    특이국소 sync-export 와 동일한 단순 요청 규격(시크릿은 body.secret,
    Content-Type text/plain)으로 preflight 를 피한다. 상세는 security.md 참조.
    """
    import hmac as _hmac_mod
    from fastapi.responses import JSONResponse
    from core.utils import _check_rate_limit
    from routers.special_sites import KSA_SYNC_SECRET, _ingest_cors

    try:
        _check_rate_limit(request, "inad_sync_export", 10, 60)
    except HTTPException as e:
        return JSONResponse({"detail": e.detail}, status_code=e.status_code,
                            headers=_ingest_cors())
    try:
        data = await request.json()
    except Exception:
        return JSONResponse({"detail": "잘못된 요청 형식"}, status_code=400,
                            headers=_ingest_cors())
    secret = str(data.get("secret") or "")
    if not (KSA_SYNC_SECRET and secret
            and _hmac_mod.compare_digest(secret, KSA_SYNC_SECRET)):
        return JSONResponse({"detail": "unauthorized"}, status_code=401,
                            headers=_ingest_cors())
    try:
        year = int(data.get("year") or 0)
    except (TypeError, ValueError):
        year = 0
    if not year:
        return JSONResponse({"detail": "year 필요"}, status_code=400,
                            headers=_ingest_cors())

    def _export():
        conn = sqlite3.connect(_INSP_DB, timeout=60)
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT 허가번호, 통합시설코드, 호출명칭, 주소, skt본부, region, ons팀, "
            "검사일자, 시정기한, 불합격내용, 불합격상세, status, 심의차수, "
            "updated_by, updated_at FROM inadequate_management WHERE year=?",
            (year,)).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    items = await asyncio.to_thread(_export)
    logger.info(f"부적합 sync-export: year={year}, {len(items)}건 (kca 미러링)")
    return JSONResponse({"success": True, "year": year,
                         "items": items, "total": len(items)},
                        headers=_ingest_cors())


@router.get("/inadequate/list")
async def inadequate_list(
    request: Request,
    year: int = Query(...),
    region: str = Query(""),
    team: str = Query(""),
    status: str = Query(""),
    kind: str = Query("", description="부적합 유형: 주소 | 공용화 | 미분류"),
    search_field: str = Query(""),
    search_values: str = Query(""),
    sort_by: str = Query(""),
    sort_dir: str = Query("desc"),
    page: int = Query(1),
    pageSize: int = Query(100),
):
    """부적합 관리 목록 조회.

    유형(주소/공용화)은 저장하지 않고 불합격상세에서 매번 계산한다. 규칙을 고쳐도
    재동기화가 필요 없고, 원천이 바뀌면 자동으로 따라간다.
    kind 필터는 SQL 로 못 거르므로(파생값) 전량을 읽어 파이썬에서 거른 뒤 페이징한다.
    2026년 대상이 58건이라 비용이 없다.
    """
    await _verify_auth(request)

    def _list():
        conn = sqlite3.connect(_INSP_DB, timeout=60)
        conn.row_factory = sqlite3.Row
        where = "year=?"
        params: list = [year]
        if region:
            where += " AND region=?"
            params.append(region)
        if team:
            where += " AND ons팀=?"
            params.append(team)
        if status:
            where += " AND status=?"
            params.append(status)
        if search_field and search_values:
            tokens = [t.strip() for t in search_values.split(',') if t.strip()]
            if tokens:
                if search_field == 'license':
                    placeholders = ','.join('?' * len(tokens))
                    normalized = [t.replace('-', '') for t in tokens]
                    where += f" AND REPLACE(허가번호, '-', '') IN ({placeholders})"
                    params.extend(normalized)
                elif search_field == 'callname':
                    clauses = ' OR '.join(['호출명칭 LIKE ?' for _ in tokens])
                    where += f" AND ({clauses})"
                    params.extend([f'%{t}%' for t in tokens])
                elif search_field == 'address':
                    clauses = ' OR '.join(['주소 LIKE ?' for _ in tokens])
                    where += f" AND ({clauses})"
                    params.extend([f'%{t}%' for t in tokens])
        offset = (page - 1) * pageSize
        _ALLOWED_SORT = {
            'region', 'ons팀', '허가번호', '호출명칭', '주소',
            '검사일자', '시정기한', '불합격내용', '불합격상세', 'status', '심의차수'
        }
        sort_col = sort_by if sort_by in _ALLOWED_SORT else '검사일자'
        dir_kw = 'ASC' if sort_dir.lower() == 'asc' else 'DESC'
        null_last = f'CASE WHEN "{sort_col}" IS NULL OR "{sort_col}" = \'\' THEN 1 ELSE 0 END'
        rows = conn.execute(
            f"SELECT * FROM inadequate_management WHERE {where} "
            f"ORDER BY {null_last}, \"{sort_col}\" {dir_kw}",
            params,
        ).fetchall()
        conn.close()

        items = []
        counts = {"주소": 0, "공용화": 0, "미분류": 0}
        for r in rows:
            d = dict(r)
            d["부적합유형"] = _inad_kind(d.get("불합격상세"))
            counts[d["부적합유형"]] += 1
            items.append(d)
        if kind:
            items = [d for d in items if d["부적합유형"] == kind]
        # 유형별 건수는 필터와 무관하게 전체 기준으로 돌려준다 — 드롭다운에
        # '주소(12)'처럼 붙여 고르기 전에 규모를 알 수 있게 한다.
        return {"items": items[offset:offset + pageSize],
                "total": len(items), "kind_counts": counts}

    return await asyncio.to_thread(_list)


@router.put("/inadequate/update")
async def inadequate_update(request: Request, req: InadequateUpdateReq):
    """부적합 상태/심의차수 업데이트 (관리자/매니저)."""
    empno = await _verify_auth(request)
    role = await asyncio.to_thread(_get_user_role_sync, empno)
    if role not in {"admin", "manager"}:
        raise HTTPException(403, "관리자/매니저만 가능")
    now = datetime.now(timezone.utc).isoformat()

    def _update():
        conn = sqlite3.connect(_INSP_DB, timeout=60)
        sets = []
        params = []
        if req.status:
            sets.append("status=?")
            params.append(req.status)
        if req.심의차수 is not None:
            sets.append("심의차수=?")
            params.append(req.심의차수)
        sets.append("updated_by=?")
        params.append(empno)
        sets.append("updated_at=?")
        params.append(now)
        params.append(req.id)
        conn.execute(f"UPDATE inadequate_management SET {','.join(sets)} WHERE id=?", params)
        conn.commit()
        conn.close()

    await asyncio.to_thread(_update)
    return {"success": True}


@router.get("/inadequate/stats")
async def inadequate_stats(
    request: Request,
    year: int = Query(...),
    region: str = Query(""),
    team: str = Query(""),
):
    """부적합 관리 통계 (필터 적용)."""
    await _verify_auth(request)

    def _stats():
        conn = sqlite3.connect(_INSP_DB, timeout=60)
        where = "year=?"
        params: list = [year]
        if region:
            where += " AND region=?"
            params.append(region)
        if team:
            where += " AND ons팀=?"
            params.append(team)
        total = conn.execute(f"SELECT COUNT(*) FROM inadequate_management WHERE {where}", params).fetchone()[0]
        done = conn.execute(f"SELECT COUNT(*) FROM inadequate_management WHERE {where} AND status='완료'", params).fetchone()[0]
        pending = conn.execute(f"SELECT COUNT(*) FROM inadequate_management WHERE {where} AND status='미완료'", params).fetchone()[0]
        excluded = conn.execute(f"SELECT COUNT(*) FROM inadequate_management WHERE {where} AND status='대상제외'", params).fetchone()[0]
        conn.close()
        return {"total": total, "완료": done, "미완료": pending, "대상제외": excluded}

    return await asyncio.to_thread(_stats)


@router.get("/inadequate/export-xlsx")
async def inadequate_export_xlsx(
    request: Request,
    year: int = Query(...),
    region: str = Query(""),
    team: str = Query(""),
    status: str = Query(""),
):
    """부적합 관리 Excel 내보내기."""
    await _verify_auth(request)
    try:
        import openpyxl
        HAS_OPENPYXL = True
    except ImportError:
        HAS_OPENPYXL = False
    if not HAS_OPENPYXL:
        raise HTTPException(503, "openpyxl 미설치")

    def _build():
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        conn = sqlite3.connect(_INSP_DB, timeout=60)
        conn.row_factory = sqlite3.Row
        where = "year=?"
        params: list = [year]
        if region:
            where += " AND region=?"
            params.append(region)
        if team:
            where += " AND ons팀=?"
            params.append(team)
        if status:
            where += " AND status=?"
            params.append(status)
        rows = conn.execute(
            f"SELECT * FROM inadequate_management WHERE {where} ORDER BY 검사일자 DESC",
            params,
        ).fetchall()
        conn.close()

        if not rows:
            raise ValueError("조회된 데이터가 없습니다")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "부적합관리"

        _thin_side = Side(style='thin')
        _thin_border = Border(left=_thin_side, right=_thin_side,
                              top=_thin_side, bottom=_thin_side)
        _hdr_fill = PatternFill('solid', fgColor='FFBFBFBF')
        _hdr_font = Font(name='맑은 고딕', size=10, bold=True)
        _data_font = Font(name='맑은 고딕', size=10)
        _center = Alignment(horizontal='center', vertical='center', wrap_text=False)
        _left = Alignment(horizontal='left', vertical='center', wrap_text=False)
        _left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
        _hdr_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
        _LINE_H = 16.5

        headers = ['본부', '팀', '허가번호', '호출명칭', '주소', '검사일자',
                   '시정기한', '불합격내용', '불합격상세', '상태', '심의차수', '최종수정자', '최종수정일시']
        db_cols = ['region', 'ons팀', '허가번호', '호출명칭', '주소', '검사일자',
                   '시정기한', '불합격내용', '불합격상세', 'status', '심의차수', 'updated_by', 'updated_at']
        _left_cols = {5, 9}
        _left_wrap_cols = {5, 9}

        ws.row_dimensions[1].height = _LINE_H
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = _hdr_font
            cell.fill = _hdr_fill
            cell.border = _thin_border
            cell.alignment = _hdr_wrap

        def _col_width(s):
            w = 0.0
            for ch in str(s):
                w += 2.2 if ord(ch) > 127 else 1.1
            return w

        all_row_values = []
        for ri, row in enumerate(rows, 2):
            d = dict(row)
            values = [d.get(col) or '' for col in db_cols]
            all_row_values.append(values)
            for ci, v in enumerate(values, 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.font = _data_font
                cell.border = _thin_border
                cell.alignment = _left_wrap if ci in _left_wrap_cols else (_left if ci in _left_cols else _center)

        col_widths = {}
        for ci in range(1, len(headers) + 1):
            best = _col_width(headers[ci - 1])
            for ri2 in range(2, len(rows) + 2):
                val = ws.cell(row=ri2, column=ci).value
                if val is not None:
                    for line in str(val).split('\n'):
                        best = max(best, _col_width(line))
            max_w = 40 if ci in _left_wrap_cols else 60
            final_w = max(min(best + 1, max_w), 10)
            col_widths[ci] = final_w
            ws.column_dimensions[get_column_letter(ci)].width = final_w

        for ri, values in enumerate(all_row_values, 2):
            max_lines = 1
            for ci, v in enumerate(values, 1):
                if not isinstance(v, str) or not v:
                    continue
                col_w = col_widths.get(ci, 10)
                cell_lines = 0
                for segment in v.split('\n'):
                    if ci in _left_wrap_cols and col_w > 0:
                        seg_w = _col_width(segment)
                        cell_lines += max(1, int(seg_w / col_w) + (1 if seg_w % col_w > 0 else 0))
                    else:
                        cell_lines += 1
                max_lines = max(max_lines, cell_lines)
            ws.row_dimensions[ri].height = _LINE_H * max_lines

        ws.freeze_panes = 'C2'
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        buf.seek(0)
        return buf.getvalue()

    try:
        data = await asyncio.to_thread(_build)
    except ValueError as e:
        logger.warning(f"부적합 xlsx 빌드 실패: {e}")
        raise HTTPException(404, "조회된 데이터가 없습니다")

    suffix_parts = [p for p in [region, team, str(year)] if p]
    filename = f"부적합관리_{'_'.join(suffix_parts)}.xlsx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
